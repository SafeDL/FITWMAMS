"""Build the full-framework disclosure from its local Markdown source."""
from pathlib import Path
import re
import sys

sys.path.insert(0, '/tmp/patent_drafting_deps')
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx_helpers import math as native_math

OUT = Path(__file__).resolve().parents[1]
SOURCE = OUT / '技术交底书_面向多失效域的自动驾驶交互测试与修复方法.md'
TARGET = SOURCE.with_suffix('.docx')


def text_run(p, s, size=11, bold=False):
    r = p.add_run(s)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.bold = bold
    fonts = r._element.get_or_add_rPr().rFonts
    fonts.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
    fonts.set(qn('w:ascii'), 'Times New Roman')
    fonts.set(qn('w:hAnsi'), 'Times New Roman')
    return r


def rich(p, s, size=11, bold=False):
    # All mathematical notation is inserted as genuine Office Math objects.
    for part in re.split(r'(\$[^$]+\$|\*\*[^*]+\*\*)', s):
        if part.startswith('$') and part.endswith('$'):
            p._p.append(native_math(part[1:-1]))
        elif part.startswith('**') and part.endswith('**'):
            text_run(p, part[2:-2], size, True)
        elif part:
            text_run(p, part, size, bold)


def paragraph(s, size=11, indent=True):
    p = D.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(.74 if indent else 0)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    rich(p, s, size)
    return p


def table(lines, cover=False):
    rows = [[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', c) for c in row)]
    t = D.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [3.8, 12.5] if cover else [3.1, 7.0, 6.2]
    assert len(widths) == len(rows[0])
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for n, (c, value) in enumerate(zip(cells, row)):
            c.width = Cm(widths[n])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            rich(p, value, 10 if cover else 9.5, index == 0)
            if index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'EEEEEE')
                c._tc.get_or_add_tcPr().append(shade)
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            trpr.append(OxmlElement('w:tblHeader'))
    paragraph('', indent=False).paragraph_format.space_after = Pt(0)


def equation(s):
    match = re.fullmatch(r'\$\$(.*?)\s+\\qquad\s+\((\d+)\)\$\$', s)
    assert match, s
    tex, number = match.groups()
    p = D.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(6), Pt(6)
    pf.keep_with_next = True
    pf.tab_stops.add_tab_stop(Cm(8.15), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(16.3), WD_TAB_ALIGNMENT.RIGHT)
    text_run(p, '\t')
    p._p.append(native_math(tex))
    text_run(p, '\t（' + number + '）', 10)


def hyperlink(p, label, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Times New Roman')
    fonts.set(qn('w:hAnsi'), 'Times New Roman')
    pr.append(fonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '18')
    pr.append(sz)
    r.append(pr)
    text = OxmlElement('w:t')
    text.text = label
    r.append(text)
    link.append(r)
    p._p.append(link)


D = Document()
sec = D.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.35)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.1)
sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
normal = D.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11)
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
for name in ['Heading 1', 'Heading 2']:
    D.styles[name].font.color.rgb = RGBColor(0, 0, 0)
settings = D.settings.element
compat = settings.find(qn('w:compat'))
if compat is None:
    compat = OxmlElement('w:compat')
    settings.append(compat)
for el in list(compat):
    if el.get(qn('w:name')) == 'compatibilityMode':
        compat.remove(el)
mode = OxmlElement('w:compatSetting')
for key, value in [('name', 'compatibilityMode'), ('uri', 'http://schemas.microsoft.com/office/word'), ('val', '15')]:
    mode.set(qn('w:' + key), value)
compat.append(mode)
D.core_properties.title = '一种面向多失效域的自动驾驶交互测试与修复方法'
D.core_properties.subject = '技术交底书任务方法论修订稿'
D.core_properties.author = '蒋拯民'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
text_run(footer, '— ', 9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
text_run(footer, ' —', 9)

lines = SOURCE.read_text().splitlines()
i = 0
while i < len(lines):
    s = lines[i].strip()
    if not s:
        i += 1
        continue
    if s.startswith('|'):
        group = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            group.append(lines[i].strip())
            i += 1
        table(group, cover=len(group[0].split('|')) == 4)
        continue
    if s.startswith('# '):
        p = paragraph(s[2:], size=18, indent=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(15)
        for r in p.runs:
            r.bold = True
    elif s.startswith('## '):
        p = D.add_paragraph(style='Heading 1')
        rich(p, s[3:], 13, True)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(12), Pt(7)
        p.paragraph_format.keep_with_next = True
    elif s.startswith('### '):
        p = D.add_paragraph(style='Heading 2')
        rich(p, s[4:], 11.5, True)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(5)
        p.paragraph_format.keep_with_next = True
    elif s.startswith('$$'):
        equation(s)
    elif s.startswith('!['):
        m = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', s)
        p = D.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(5), Pt(3)
        p.paragraph_format.keep_with_next = True
        picture = p.add_run().add_picture(str((OUT / m.group(2)).resolve()), width=Cm(16.1))
        picture._inline.docPr.set('descr', m.group(1))
    elif s.startswith('**图'):
        p = paragraph(s[2:-2], indent=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = '附图占位' in s
        for r in p.runs:
            r.bold = True
    elif s.startswith('[') and re.match(r'\[\d+\]', s):
        reference_url = re.search(r'https?://\S+$', s)
        assert reference_url, s
        before, url = s[:reference_url.start()].rstrip(), reference_url.group(0)
        p = paragraph(before, size=9, indent=False)
        p.paragraph_format.line_spacing = 1.1
        p.paragraph_format.space_after = Pt(4)
        text_run(p, ' ', 9)
        hyperlink(p, '原文链接', url)
    else:
        p = paragraph(s)
        following = next((line.strip() for line in lines[i+1:] if line.strip()), '')
        if following.startswith('$$'):
            p.paragraph_format.keep_with_next = True
    i += 1

D.save(TARGET)
print(TARGET)
