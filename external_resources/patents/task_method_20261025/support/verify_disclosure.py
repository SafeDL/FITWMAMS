"""Check the full-framework DOCX, source preservation and numerical evidence."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import sys
from lxml import etree

sys.path.insert(0, "/tmp/patent_drafting_deps")
from docx import Document

ROOT = Path(__file__).resolve().parents[4]
OUT = Path(__file__).resolve().parents[1]
MD = OUT / "技术交底书_面向多失效域的自动驾驶交互测试与修复方法.md"
DOC = MD.with_suffix(".docx")
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}


def read_json(path):
    return json.loads((ROOT / path).read_text())


source = MD.read_text()
inline = len(re.findall(r"(?<!\$)\$[^$]+\$(?!\$)", source))
display = len(re.findall(r"^\$\$", source, re.M))
assert display == 20
with ZipFile(DOC) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            etree.fromstring(z.read(name))
    xml = etree.fromstring(z.read("word/document.xml"))
    settings = etree.fromstring(z.read("word/settings.xml"))
    equations = xml.findall(".//m:oMath", NS)
    assert len(equations) == inline + display
    for equation in equations:
        text = "".join(equation.xpath(".//m:t/text()", namespaces=NS))
        assert not re.search(r"[\u4e00-\u9fff，。；：]", text), text
    assert not xml.xpath(
        "//w:pict | //w:object | //w:commentRangeStart | //w:ins | //w:del",
        namespaces=NS,
    )
    assert not xml.xpath('//w:instrText[contains(.,"EQ ")]', namespaces=NS)
    normal = "".join(xml.xpath("//w:t/text()", namespaces=NS))
    assert not re.search(r"\$|\\(?:frac|omega|tau)|[_^]", normal)
    numbers = [
        s for s in xml.xpath("//w:t/text()", namespaces=NS)
        if re.fullmatch(r"（\d+）", s)
    ]
    assert numbers == [f"（{i}）" for i in range(1, 21)]
    assert settings.xpath(
        '//w:compatSetting[@w:name="compatibilityMode"]/@w:val',
        namespaces=NS,
    ) == ["15"]
    headings = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading1"]]', namespaces=NS)
    subheadings = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading2"]]', namespaces=NS)
    assert len(headings) == 7 and len(subheadings) == 6
    assert normal.count("（附图占位）") == 4
    media = sorted(n for n in z.namelist() if n.startswith("word/media/"))
    assert len(media) == 2 and all(n.endswith(".png") for n in media)
    image_paths = [
        (OUT / path).resolve()
        for path in re.findall(r"^!\[.*?\]\((.*?)\)$", source, re.M)
    ]
    assert sorted(hashlib.sha256(z.read(n)).hexdigest() for n in media) == sorted(
        hashlib.sha256(p.read_bytes()).hexdigest() for p in image_paths
    )
    captions = re.findall(r"^\*\*图(\d+) ", source, re.M)
    assert captions == [str(i) for i in range(1, 7)]
    refs = re.findall(r"^\[(\d+)\]", source, re.M)
    assert refs == [str(i) for i in range(1, 9)]
    cited = set()
    for group in re.findall(r"\[(\d+(?:,\d+)*)\]", source.split("## 7 ")[0]):
        cited.update(group.split(","))
    assert cited == set(refs)
    expected_urls = [
        re.search(r"https?://\S+$", s).group(0)
        for s in source.splitlines() if re.match(r"^\[\d+\]", s)
    ]
    rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
    actual_urls = [
        r.get("Target") for r in rels if r.get("Type").endswith("/hyperlink")
    ]
    assert actual_urls == expected_urls
    for line in source.splitlines():
        if re.match(r"^\[\d+\]", line):
            prefix = line[:re.search(r"https?://\S+$", line).start()].rstrip()
            assert prefix in normal
    for term in ["极值理论", "归一化流", "条件扩散", "计划适用性", "分支", "纠正", "持续失效"]:
        assert term in normal
    assert "世界生成及执行参数" not in normal or "保持固定" in normal
    assert "有限性能关系检查形成经验搜索分支" in normal
    assert "完整执行及未停车资格" in normal
    for term in ["已实际执行偏离参考的控制", "本层执行前固定", "按其当前成员关系", "相同执行限幅处理", "先选择分支再选择该分支的输入组", "条件采样目标", "详细平衡", "投影后场景", "经验比例"]:
        assert term in normal, term

readback = Document(DOC)
assert readback.core_properties.title == "一种面向多失效域的自动驾驶交互测试与修复方法"
assert len(readback.tables) == 2
assert not list(OUT.rglob("*.pdf")) and not list(OUT.rglob("*.svg"))
abstract = source.split("**摘要建议稿**")[1].strip()
assert len(abstract) < 300

claims = (OUT / "候选权利要求与说明书支持.md").read_text()
assert re.findall(r"^\*\*(\d+)\.\*\*", claims, re.M) == [str(i) for i in range(1, 16)]
assert "权利要求1至13" in claims
assert "或者用于当前策略重新交互" in claims

manifest = json.loads((OUT / "support/source_manifest.json").read_text())
assert len(manifest) == 58
for item in manifest:
    assert hashlib.sha256((ROOT / item["path"]).read_bytes()).hexdigest() == item["sha256"], item["path"]

evt = read_json("results/evt/evt_summary.json")
assert evt["source"]["all_segments"] == 96055 and evt["num_tail_segments"] == 2964
flow = read_json("results/flow/evaluation_summary.json")
assert flow["final_fit_events"] == 2964
for name, expected in [("flow", 0.0615), ("gaussian", 0.1332), ("gmm8", 0.1046)]:
    assert round(flow["method_comparison"][name]["sliced_wasserstein_256"], 4) == expected
diffusion = read_json("results/diffusion/evaluation_summary.json")
assert round(diffusion["headline_metrics"]["mean_single_draw_ADE_m"], 4) == 0.0527
assert diffusion["condition_disclosure"]["uses_future_background_position_knots"]
assert not diffusion["condition_disclosure"]["uses_future_ego"]
bridge = read_json("results/hierarchical_world_model/bridge/test_report.json")
assert bridge["unique_scenes"] == 10151
for key in ["takeover_20261015", "takeover_20261016", "takeover_20261017"]:
    assert bridge["summary"][key]["admitted_scenes"] == 7
interventions = read_json("results/hierarchical_world_model/bridge/interventions.json")
assert interventions["cases"] == 51 and interventions["runs"] == 204
braking = interventions["summary"]["braking"]
assert braking["cases"] == 9 and braking["bridge_runs"] == 27
assert braking["resolved_collision_pairs"] == 15 and braking["introduced_collision_pairs"] == 0
assert not interventions["real_ADS_policy"] and not interventions["parameters_fitted_on_these_cases"]
latency = read_json("results/hierarchical_world_model/bridge/latency.json")
assert sum(row["steps"] for row in latency["cpu"].values()) == 8493
assert round(latency["P99_ms"], 2) == 8.59 and not latency["ADS_compute_included"]
base = "results/risk_stratified_repair/formal/diffusion_planner/preparation/"
search = read_json(base + "round_one_search_comparison.json")
for method, expected in [("monte_carlo", (160, 158)), ("subset", (215, 212)), ("guided", (450, 449))]:
    row = search["methods"][method]
    assert row["executions"] == 5000
    assert (row["primary_unique_collision_instances"], row["collisions_after_safe_initial_state"]) == expected
assert search["methods"]["guided"]["collision_sampling_roots"] == 98
assert search["methods"]["guided"]["diagnostic_collision_calls"] == 19
assert search["methods"]["guided"]["archive_failure_representatives"] == 106
assert search["methods"]["guided"]["archive_success_neighbours"] == 9
replay = read_json(base + "guided_round_two_replay_acceptance.json")
assert replay["archive_endpoints"] == 115
assert replay["all_endpoints"]["safety_repairs"] == 43
assert replay["all_endpoints"]["original_safety_failures"] - 43 == 63
assert replay["old_success_endpoints"] == 9
assert replay["old_success_regressions"] == 1
assert "current search branches" in replay["feedback_scope"]
materials = read_json(base + "round_one_material_acceptance.json")["methods"]["guided"]
assert materials["expert_inputs"] == 27 and materials["policy_success_inputs"] == 37
admission = read_json(base + "guided_round_one_admission_acceptance.json")
assert admission["selected_candidate_updates"] == 500
assert admission["full_checkpoint_payload_equal_to_selected"]
for candidate, expected in zip(admission["candidates"], [(15, 9), (50, 27)]):
    comparisons = candidate["comparisons"]
    actual = tuple(
        round(comparisons["official/" + kind]["collision_rate_after"] * 500)
        for kind in ["tail_generated", "natural_generated"]
    )
    assert actual == expected
completed = read_json(base + "completed_delivery_acceptance.json")
assert not completed["nuplan_city_simulation_performed"]
assert not completed["model_selection_or_training_from_final_results"]
assert "61 to 40" in completed["conclusions"]["tail_guided_vs_initial"]
assert "76 to 80" in completed["conclusions"]["natural_empirical_guided_vs_initial"]
assert "41 to 37" in completed["conclusions"]["natural_generated_guided_vs_initial"]

report = {
    "date": "2026-10-25",
    "scope": "task_focused_method_formulation_and_prior_art",
    "structural_checks_passed": True,
    "native_equations": len(equations),
    "inline_equations": inline,
    "numbered_equations": display,
    "main_sections": len(headings),
    "tables": len(readback.tables),
    "implementation_subsections": len(subheadings),
    "existing_png_images": len(media),
    "logic_figure_placeholders": 4,
    "new_images_generated": 0,
    "claims_checked": 15,
    "reference_links_checked": len(actual_urls),
    "sources_hash_checked": len(manifest),
    "terminology_registry": "support/术语与符号核对.md",
    "numerical_evidence_checked": True,
    "new_training_simulation_or_expert_calls": 0,
    "abstract_characters": len(abstract),
    "docx_sha256": hashlib.sha256(DOC.read_bytes()).hexdigest(),
    "word_desktop_reopen": False,
    "word_note": "No desktop Word session. Native Office Math, ZIP/XML and python-docx readback checked; LibreOffice rendering used for layout.",
    "pdf_delivered": False,
    "svg_delivered": False,
}
report_path = OUT / "support/qa_report.json"
if report_path.exists():
    previous = json.loads(report_path.read_text())
    if previous.get("docx_sha256") == report["docx_sha256"]:
        for key in ["layout_pages", "page_limit", "within_page_limit", "layout_review", "equation_number_alignment", "page_bounds_checked", "reference_order", "approved_previous_version_preserved"]:
            if key in previous:
                report[key] = previous[key]
report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
if "layout_pages" in report:
    assert report["layout_pages"] <= 20
print(json.dumps(report, ensure_ascii=False))
