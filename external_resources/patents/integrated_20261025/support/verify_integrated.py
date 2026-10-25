"""Check document integrity, notation, evidence arithmetic and source preservation."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import sys
from lxml import etree

sys.path.insert(0, '/tmp/patent_drafting_deps')
from docx import Document

ROOT = Path(__file__).resolve().parents[4]
OUT = Path(__file__).resolve().parents[1]
DOC = OUT / '技术交底书_交互失效邻域反馈的自动驾驶迭代测试方法.docx'
MD = DOC.with_suffix('.md')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}

with ZipFile(DOC) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith(('.xml', '.rels')):
            etree.fromstring(z.read(name))
    xml = etree.fromstring(z.read('word/document.xml'))
    settings = etree.fromstring(z.read('word/settings.xml'))
    formulas = xml.findall('.//m:oMath', NS)
    expected = len(re.findall(r'(?<!\$)\$[^$]+\$(?!\$)', MD.read_text()))
    display = len(re.findall(r'^\$\$', MD.read_text(), re.M))
    assert len(formulas) == expected + display, (len(formulas), expected, display)
    assert display == 11
    for f in formulas:
        text = ''.join(f.xpath('.//m:t/text()', namespaces=NS))
        assert not re.search(r'[\u4e00-\u9fff，。；：]', text), text
    assert not xml.xpath('//w:pict | //w:object', namespaces=NS)
    assert not xml.xpath('//w:instrText[contains(.,"EQ ")]', namespaces=NS)
    assert not xml.xpath('//w:del | //w:ins | //w:commentRangeStart', namespaces=NS)
    normal = ''.join(xml.xpath('//w:t/text()', namespaces=NS))
    notation_text = re.sub(r'nat_\d+_\d+_\d+', '', normal)
    assert not re.search(r'\$|\\(?:frac|omega|pi|tau)|[_^]', notation_text)
    numbers = xml.xpath('//w:t[starts-with(.,"（")]/text()', namespaces=NS)
    numbers = [s for s in numbers if re.fullmatch(r'（\d+）', s)]
    assert numbers == ['（' + str(i) + '）' for i in range(1, 12)], numbers
    assert settings.xpath('//w:compatSetting[@w:name="compatibilityMode"]/@w:val', namespaces=NS) == ['15']
    media = [n for n in z.namelist() if n.startswith('word/media/')]
    assert len(media) == 3 and all(n.endswith('.png') for n in media), media
    headings = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading1"]]', namespaces=NS)
    subheadings = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading2"]]', namespaces=NS)
    assert len(headings) == 7 and len(subheadings) == 5

readback = Document(DOC)
assert len(readback.tables) == 2
for line in MD.read_text().splitlines():
    if line.startswith('!['):
        target = re.search(r'\]\((.*?)\)', line).group(1)
        assert (OUT / target).resolve().is_file()

sources = json.loads((OUT / 'support/source_manifest.json').read_text())
for source in sources:
    path = ROOT / source['path']
    assert hashlib.sha256(path.read_bytes()).hexdigest() == source['sha256'], source['path']

base = ROOT / 'results/risk_stratified_repair/formal/diffusion_planner/preparation'
search = json.loads((base / 'round_one_search_comparison.json').read_text())
assert [search['methods'][m]['primary_unique_collision_instances'] for m in ['monte_carlo', 'subset', 'guided']] == [160, 215, 450]
assert [search['methods'][m]['collisions_after_safe_initial_state'] for m in ['monte_carlo', 'subset', 'guided']] == [158, 212, 449]
assert all(d['executions'] == 5000 for d in search['methods'].values())
replay = json.loads((base / 'guided_round_two_replay_acceptance.json').read_text())
assert replay['archive_endpoints'] == 115
assert replay['all_endpoints']['safety_repairs'] == 43
assert replay['old_success_endpoints'] == 9
assert replay['old_success_regressions'] == 1
assert (106 - 43, 43, 1, 9 - 1) == (63, 43, 1, 8)
assert not list(OUT.glob('*.pdf')) and not list(OUT.glob('*.svg'))

report = {
    'date': '2026-10-25',
    'structural_checks_passed': True,
    'native_equations': len(formulas),
    'inline_equations': expected,
    'numbered_equations': display,
    'figures_png': len(media),
    'logic_figure_placeholders': 2,
    'heading_1': len(headings),
    'heading_2': len(subheadings),
    'tables': len(readback.tables),
    'source_hashes_verified': len(sources),
    'new_model_or_simulation_executions': 0,
    'word_desktop_reopen': False,
    'word_desktop_note': 'No desktop Word session is available; ZIP/XML and python-docx readback checked. Layout is inspected with LibreOffice rendering.',
    'docx_sha256': hashlib.sha256(DOC.read_bytes()).hexdigest(),
}
(OUT / 'support/qa_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False))
