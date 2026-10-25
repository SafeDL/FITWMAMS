from pathlib import Path
from zipfile import ZipFile
import hashlib,json,re,subprocess,sys
from lxml import etree
sys.path.insert(0,'/tmp/patent_drafting_deps')
from docx import Document
O=Path(__file__).resolve().parents[1]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
reports={}
for p in O.glob('*.docx'):
 with ZipFile(p) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith('.xml'):etree.fromstring(z.read(name))
  r=etree.fromstring(z.read('word/document.xml'));st=etree.fromstring(z.read('word/settings.xml'))
  wt=r.xpath('//w:t/text()',namespaces=ns);mt=r.xpath('//m:t/text()',namespaces=ns)
  vals={
   'zip_xml_valid':True,'docx_loader_reopened':bool(Document(p)),
   'math_objects':len(r.findall('.//m:oMath',ns)),
   'display_equations':len(r.findall('.//m:oMathPara',ns)),
   'tables':len(r.findall('.//w:tbl',ns)),
   'legacy_pict_objects':len(r.findall('.//w:pict',ns)),
   'old_equation_fields':len([x for x in r.xpath('//w:instrText/text()',namespaces=ns) if re.search(r'EQ|EMBED',x)]),
   'equationxml_parts':len([n for n in z.namelist() if 'equationxml' in n.lower()]),
   'media_objects':len([n for n in z.namelist() if n.startswith('word/media/')]),
   'chinese_or_punctuation_in_equations':[x for x in mt if re.search('[\u4e00-\u9fff，。；：]',x)],
   'unrendered_math_markers':[x for x in wt if '$' in x or re.search(r'\\(?:frac|sum|sqrt|theta|omega|pi)',x)],
   'compatibility_mode':st.xpath('string(//w:compatSetting[@w:name="compatibilityMode"]/@w:val)',namespaces=ns),
   'comments_parts':[n for n in z.namelist() if 'comments' in n],
   'tracked_insertions':len(r.findall('.//w:ins',ns)),
   'tracked_deletions':len(r.findall('.//w:del',ns)),
   'text_characters':len(''.join(wt)),
   'sha256':hashlib.sha256(p.read_bytes()).hexdigest()
  }
  assert vals['legacy_pict_objects']==vals['old_equation_fields']==vals['equationxml_parts']==vals['media_objects']==0
  assert not vals['chinese_or_punctuation_in_equations'] and not vals['unrendered_math_markers']
  assert vals['compatibility_mode']=='15'
  if '修订稿' in p.name:
   assert vals['display_equations']==14
   assert ''.join(wt).count('【附图占位：待绘制】')==4
   vals['figure_placeholders']=4
  for e in r.findall('.//m:r',ns):
   tags=[etree.QName(c).localname for c in e]
   assert tags==['rPr','rPr','t'],tags
 pdf=p.with_suffix('.pdf');assert pdf.exists()
 txt=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
 assert not re.search('[¿❑�]',txt)
 pages=txt.split('\f');assert all(len(x.strip())>30 for x in pages if x.strip())
 info=subprocess.check_output(['pdfinfo',str(pdf)],text=True)
 vals['pdf_pages']=int(re.search(r'Pages:\s+(\d+)',info).group(1))
 vals['pdf_render_no_error_placeholders']=True
 reports[p.name]=vals
result={'date':'2026-10-23','checks_passed':True,'documents':reports,'visual_review':'All-page contact sheets inspected; larger views checked for equation blocks 6–11. No clipped formula or image formula.','office_validation':'LibreOffice 7.3 DOCX load and PDF export passed. Microsoft Word desktop is unavailable; Word-native reopening was not performed.','sources_unchanged_by_task':True,'experiments_started':False}
(O/'support/qa_report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
