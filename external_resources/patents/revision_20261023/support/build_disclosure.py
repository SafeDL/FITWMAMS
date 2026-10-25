from pathlib import Path
import sys,re,json,hashlib
sys.path.insert(0,'/tmp/patent_drafting_deps')
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree
import latex2mathml.converter

OUT=Path(__file__).resolve().parents[1]
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def font(run,size=11,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    rpr=run._element.get_or_add_rPr();fs=rpr.rFonts
    fs.set(qn('w:eastAsia'),'Noto Serif CJK SC');fs.set(qn('w:ascii'),'Times New Roman');fs.set(qn('w:hAnsi'),'Times New Roman')

def math(tex):
    """Convert the small presentation-MathML vocabulary used here to native OMML."""
    root=etree.fromstring(latex2mathml.converter.convert(tex).encode())
    def E(tag,**attrs):
        e=etree.Element('{'+M+'}'+tag)
        for k,v in attrs.items():e.set('{'+M+'}'+k,str(v))
        return e
    def run(txt,style='i',variant=''):
        e=E('r');pr=E('rPr')
        if variant in ('script','double-struck','fraktur'):
            pr.append(E('scr',val=variant));pr.append(E('sty',val='p'))
        else:pr.append(E('sty',val='b' if variant=='bold' else style))
        e.append(pr)
        wpr=etree.SubElement(e,'{'+W+'}rPr')
        fs=etree.SubElement(wpr,'{'+W+'}rFonts')
        fs.set('{'+W+'}ascii','Cambria Math');fs.set('{'+W+'}hAnsi','Cambria Math')
        sz=etree.SubElement(wpr,'{'+W+'}sz');sz.set('{'+W+'}val','22')
        tx=E('t');tx.text=txt
        if txt.startswith(' ') or txt.endswith(' '):tx.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        e.append(tx);return e
    def container(tag,items):
        e=E(tag)
        for x in items:e.append(x)
        return e
    def conv(n,variant=''):
        tag=etree.QName(n).localname
        variant=n.get('mathvariant',variant)
        cs=list(n)
        if tag in ('math','mrow','mstyle','semantics','mpadded'):
            def delim(left,right,children):
                e=E('d');pr=E('dPr');pr.append(E('begChr',val=left));pr.append(E('endChr',val=right));e.append(pr)
                inner=[]
                for c in children:inner+=conv(c,variant)
                e.append(container('e',inner));return e
            if len(cs)>1 and etree.QName(cs[0]).localname=='mo' and etree.QName(cs[-1]).localname=='mo':
                left,right=cs[0].text,cs[-1].text
                if (left,right) in [('(',')'),('[',']'),('{','}'),('|','|')]:
                    return [delim(left,right,cs[1:-1])]
            out=[];idx=0
            while idx<len(cs):
                c=cs[idx]
                if etree.QName(c).localname=='mo' and c.text=='|':
                    end=next((j for j in range(idx+1,len(cs)) if etree.QName(cs[j]).localname=='mo' and cs[j].text=='|'),None)
                    if end is not None:
                        out.append(delim('|','|',cs[idx+1:end]));idx=end+1;continue
                out+=conv(c,variant);idx+=1
            return out
        if tag in ('mi','mn','mo','mtext','ms'):
            return [run('∼' if n.text=='~' else (n.text or ''), 'i' if tag=='mi' and variant!='normal' and n.text!='~' else 'p', variant)]
        if tag=='mspace':return [run(' ','p')]
        if tag in ('msub','msup','msubsup'):
            e=E({'msub':'sSub','msup':'sSup','msubsup':'sSubSup'}[tag]);e.append(container('e',conv(cs[0],variant)))
            if tag in ('msub','msubsup'):e.append(container('sub',conv(cs[1],variant)))
            if tag in ('msup','msubsup'):e.append(container('sup',conv(cs[-1],variant)))
            return [e]
        if tag=='mfrac':
            e=E('f');e.append(container('num',conv(cs[0],variant)));e.append(container('den',conv(cs[1],variant)));return [e]
        if tag in ('msqrt','mroot'):
            e=E('rad');pr=E('radPr');pr.append(E('degHide',val=1 if tag=='msqrt' else 0));e.append(pr)
            e.append(container('deg',[] if tag=='msqrt' else conv(cs[1],variant)))
            items=[]
            for c in (cs if tag=='msqrt' else cs[:1]):items+=conv(c,variant)
            e.append(container('e',items));return [e]
        if tag=='mover' and ''.join(cs[1].itertext()) in ('¯','‾','^','→','ˆ'):
            e=E('acc');pr=E('accPr');accent=''.join(cs[1].itertext())
            pr.append(E('chr',val={'¯':'\u0305','‾':'\u0305','^':'\u0302','ˆ':'\u0302','→':'\u20d7'}[accent]));e.append(pr)
            e.append(container('e',conv(cs[0],variant)));return [e]
        if tag in ('munder','mover','munderover'):
            base=conv(cs[0],variant)
            if tag in ('munder','munderover'):
                lo=E('limLow');lo.append(container('e',base));lo.append(container('lim',conv(cs[1],variant)));base=[lo]
            if tag in ('mover','munderover'):
                up=E('limUpp');up.append(container('e',base));up.append(container('lim',conv(cs[-1],variant)));base=[up]
            return base
        if tag=='annotation':return []
        raise ValueError('Unsupported MathML tag '+tag+' in '+tex)
    out=E('oMath')
    for x in conv(root):out.append(x)
    return out

def rich(p,text,size=11,bold=False):
    parts=re.split(r'(\$[^$]+\$)',text)
    for s in parts:
        if s.startswith('$') and s.endswith('$'):p._p.append(math(s[1:-1]))
        elif s:font(p.add_run(s),size,bold)

def para(d,text,style=None):
    p=d.add_paragraph(style=style);rich(p,text)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.35
    if not style:p.paragraph_format.first_line_indent=Cm(.74)
    return p

def heading(d,text,level=1):
    p=d.add_paragraph(style='Heading '+str(level));rich(p,text,14 if level==1 else 12,True)
    p.paragraph_format.space_before=Pt(12 if level==1 else 8);p.paragraph_format.space_after=Pt(6)
    return p

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid';t.autofit=False
    if widths:
        for col,width in zip(t.columns,widths):col.width=Cm(width)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i];c.text='';rich(c.paragraphs[0],h,10,True)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');c._tc.get_or_add_tcPr().append(sh)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):
            c.text='';rich(c.paragraphs[0],str(txt),10)
    for r in t.rows:
        for i,c in enumerate(r.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:c.width=Cm(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.15
        cp=OxmlElement('w:cantSplit');r._tr.get_or_add_trPr().append(cp)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def eq(d,tex,n,explain):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    mp=OxmlElement('m:oMathPara');mp.append(math(tex));p._p.append(mp)
    font(p.add_run('  （'+str(n)+'）'),10)
    para(d,explain)

def setup(title):
    d=Document();s=d.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21)
    s.top_margin=Cm(2.3);s.bottom_margin=Cm(2.2);s.left_margin=Cm(2.4);s.right_margin=Cm(2.4)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    for n in ['Heading 1','Heading 2','Heading 3']:
        d.styles[n].font.color.rgb=RGBColor(0,0,0)
    settings=d.settings.element
    compat=settings.find(qn('w:compat'))
    if compat is None:compat=OxmlElement('w:compat');settings.append(compat)
    for el in list(compat):
        if el.get(qn('w:name'))=='compatibilityMode':compat.remove(el)
    x=OxmlElement('w:compatSetting');x.set(qn('w:name'),'compatibilityMode');x.set(qn('w:uri'),'http://schemas.microsoft.com/office/word');x.set(qn('w:val'),'15');compat.append(x)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('— '),9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);font(footer.add_run(' —'),9)
    d.core_properties.title=title;d.core_properties.author='蒋拯民';d.core_properties.subject='自然驾驶事件约束的自动驾驶多失效域测试'
    return d

REFS=[
('1','Mendoza J P, Unnikrishnan R, Wahal S, Zhang L. Selectively intelligent actors in simulations. WO2025090341A1，公开日：2025-05-01。','https://patents.google.com/patent/WO2025090341A1/en'),
('2','Suo S, Wong K, Xu J, et al. MixSim: A Hierarchical Framework for Mixed Reality Traffic Simulation. CVPR，2023。','https://waabi.ai/research/mixsim'),
('3','Lin H, Chen H, Wolff E. System and method of causal composition diffusion for closed loop traffic generation. US20260134173A1，公开日：2026-05-14。','https://patents.google.com/patent/US20260134173A1/en'),
('4','Huang P, Ding W, Stoler B, et al. CaDRE: Controllable and Diverse Generation of Safety-Critical Driving Scenarios Using Real-World Trajectories. ICRA，2025。预印本：arXiv:2403.13208。','https://arxiv.org/html/2403.13208v2'),
('5','Kinnear H J, DiazDelaO F A. Niching subset simulation. Probabilistic Engineering Mechanics，2025，79：103729。DOI：10.1016/j.probengmech.2025.103729。','https://discovery.ucl.ac.uk/id/eprint/10204007/'),
('6','白先旭，李亲，李维汉，石琴。一种自动驾驶车辆在不同安全域的测试场景创建方法。CN116361175A，公开日：2023-06-30。','https://patents.google.com/patent/CN116361175A/zh'),
('7','祝亚运，单文君，余福珍。基于混合专家模型的自动驾驶测试场景生成方法和系统。CN121743096A，公开日：2026-03-27。','https://patents.google.com/patent/CN121743096A/zh'),
('8','Rempe D, Philion J, Guibas L, et al. Generating Useful Accident-Prone Driving Scenarios via a Learned Traffic Prior. CVPR，2022。','https://research.nvidia.com/labs/toronto-ai/STRIVE/'),
('9','Krajewski R, Bock J, Kloeker L, Eckstein L. The highD Dataset: A Drone Dataset of Naturalistic Vehicle Trajectories on German Highways for Validation of Highly Automated Driving Systems. ITSC，2018：2118—2125。','https://arxiv.org/abs/1810.05642'),
('10','Lin H, Huang X, Phan-Minh T, et al. Causal Composition Diffusion Model for Closed-loop Traffic Generation. CVPR，2025。预印本：arXiv:2412.17920。','https://arxiv.org/html/2412.17920v3')]

def references(d):
    for n,txt,url in REFS:
        p=para(d,'['+n+'] '+txt);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(2)
        for r in p.runs:r.font.size=Pt(9.5)
        p=para(d,url);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(4)
        for r in p.runs:r.font.size=Pt(8.5)

D=setup('技术交底书：自然驾驶事件约束的自动驾驶多失效域测试方法（修订稿）')
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;rich(p,'发明专利技术交底书',20,True)
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;rich(p,'自然驾驶事件约束的自动驾驶多失效域测试方法',15,True)
p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;rich(p,'修订稿 · 2026年10月23日',10)
table(D,['项目','内容'],[
('专利名称','一种自然驾驶事件约束的自动驾驶多失效域测试方法'),
('所属技术领域','自动驾驶仿真测试、交通行为建模与计算机辅助测试'),
('专利发明人或设计人','蒋拯民、李慧云、潘毅、颜迥超'),
('技术交底书撰写人','蒋拯民'),('技术问题联系人','蒋拯民'),('联系电话','17749028738（同微信）'),('电子邮箱／申请人','待补充')],[3.5,12.7])
heading(D,'技术方案摘要')
para(D,'本发明针对自然驾驶事件的交互扩展测试，固定道路、车辆身份、尺寸及起始状态，将后续可变行为组织为可单独替换和重放的参数组。背景车辆按参考运动执行，并在实际交互使原参考不再适用时转入响应控制，承接已开始的换道及当前运动状态。在相同事件、策略和执行规则下，依据单组参数变更前后的实际风险变化选择后续扰动；在风险准入样本中联合保留不同交互类别的代表，并跨类别限制同一独立抽样来源获得的测试名额。实际失效及完整输入另行保存，用于缺陷复现和策略改进。')
heading(D,'1 技术领域与所要解决的技术问题')
para(D,'本发明涉及自动驾驶系统的计算机仿真测试。具体地，本发明以自然驾驶记录中的多车交互事件为条件，通过可重放的行为参数调整与交互式交通执行，生成针对当前被测自动驾驶策略的失效场景，并在有限仿真预算内保留不同交互过程的测试机会。')
para(D,'自然驾驶记录描述了特定道路、参与车辆和起始状态下已经发生的一种交通过程。当被测车辆采取与记录不同的制动、加速或换道动作时，背景车辆继续按原时间轨迹运行，可能产生与实际交互不相容的运动；若不保留事件条件而整体重生成交通，又难以确定风险变化来自何种行为调整。与此同时，逐步提高风险要求的搜索容易反复扩展少数高风险样本，即使生成实例数量持续增加，也可能集中于相近的交互过程或同一抽样来源。')
para(D,'本发明所要解决的技术问题是：在保留自然驾驶事件条件并允许背景车辆响应实际交通变化的前提下，如何使行为参数调整与当前被测策略的实际风险变化建立可核对的对应关系，并将交互过程差异和样本来源集中程度共同用于后续测试安排，以形成可复现的多类失效场景集合。')
para(D,'本交底书中的“失效域”指在固定事件条件及允许行为参数空间内，由已经观察到的失效实例及其交互特征表征的经验失效区域。“交互类别”用于组织测试中的交通变化，可以包含尚未失效的高风险样本；其类别数量不直接等同于失效域数量。独立来源指一次完整独立条件抽样的谱系起点，与自然记录的录像、车辆或时间窗口来源分别标识。')
heading(D,'2 技术背景与相近的现有方案')
heading(D,'2.1 自然事件恢复与选择性响应',2)
para(D,'公开专利WO2025090341A1记载了基于日志构建仿真、在交通交互满足切换条件时使背景车辆偏离日志轨迹，并继续处理其他受影响车辆的方案，还涉及空间路径跟随及行为参数拟合。[1] MixSim利用路线目标与低层行为策略组织混合现实交通仿真，并通过相关车辆路线的优化构造危险变体。[2] 上述方案为事件保持和动态响应提供了直接技术基础。')
para(D,'本发明在此任务范围内进一步处理行为输入与测试安排的对应关系：将允许变化的行为输入细分为可替换参数组，利用同一执行机制获得的实际风险变化指导下一次参数组选择，并使风险准入后的样本保留同时受交互类别与独立来源约束。')
heading(D,'2.2 交互关系引导的危险生成',2)
para(D,'CCDiff及关联公开申请US20260134173A1根据参与者之间的交互结构对车辆排序，对关键车辆的扩散逆采样选择性施加控制目标梯度。[3][10] CN121743096A则结合场景测试结果、决策不确定性和敏感区域进行定向场景生成。[7] 因此，以交互关系或测试反馈选择干预对象，已经具有相关公开基础。')
para(D,'本发明采用的具体处理是：固定自然事件、当前策略、非选中行为输入及随机实现，从事件起点比较原实例与单组变更实例，记录变更经背景响应和被测策略作用后的实际风险差，并将该差值用于参数组选择。比较对象是可重放的环境输入，比较结果进入后续仿真输入的分配过程。')
heading(D,'2.3 多样性搜索与风险分层',2)
para(D,'STRIVE在学习到的交通先验中优化事故场景，并组织事故类型及规划器改进。[8] CaDRE将真实轨迹扰动写入质量—多样性搜索，采用行为网格档案与邻域空网格感知重启，兼顾危险程度和行为差异。[4] Niching subset simulation通过输入空间分区和递归子集模拟处理多峰失效区域。[5] CN116361175A根据自然驾驶特征分类、组合测试及风险指标，将测试用例划分至不同安全域。[6]')
para(D,'本发明着重处理风险分层过程中两种不同的集中现象：相似交互反复获得扩展机会，以及同一独立样本的派生实例占用较多测试机会。交互类别代表保留针对前者；跨类别累计并限制同一来源的名额针对后者。两种约束共同作用于同一次后续测试名额分配，实际失效的保存另行执行。')
heading(D,'3 现有方案的不足与本发明目的')
para(D,'对于上述目标任务，单独采用轨迹重放、交互生成或危险排序，尚不能直接给出“事件中保留什么、允许改变什么、风险变化如何用于后续测试”的完整实施关系。若一次提议同时改变道路、初态和多类行为输入，则其风险变化难以用于判断具体参数组的后续调整价值。若背景车辆响应后仍追赶原时间位置，则搜索还可能放大参考跟踪与交通响应之间的不相容。')
para(D,'本发明将自然驾驶事件作为固定条件，将行为参数组作为采样、替换与重放的共同接口；由统一执行模块处理参考跟踪、响应接管和车辆动力学，再以同条件比较得到的实际风险变化调整参数组选择概率，使被选择的变化能够追溯至具体输入。')
para(D,'在此基础上，本发明在风险准入集合内联合分配交互类别代表及独立来源名额，使同一来源进入不同类别后仍共同受限。由此，将测试从单个高风险样本的反复扩展，转向具有明确交互差异和来源约束的持续探索，并以完整输入档案支持后续复现。')
heading(D,'4 技术方案概述')
para(D,'本方法由包括处理器、存储器和交通仿真接口的测试设备执行。处理器调用冻结的场景生成模型、背景驾驶模型及当前被测策略，存储器保存自然事件、参数数组、执行轨迹和来源谱系。总体流程见图1。')
for x in [
'S101，构建事件条件。清洗连续自然驾驶记录，截取完整多车交互窗口，计算事件风险并选择目标事件，保存道路、车辆身份、尺寸、起始状态及有效历史。',
'S102，建立行为参数组。在固定事件条件下，分别为参考运动、驾驶员特征、换道及逐时响应设置可变输入，生成背景车辆参考运动，并为每次完整独立条件抽样建立来源标识。',
'S103，执行交互式测试。依据当前实际联合状态计算背景车辆基准动作及参考适用性；需要响应时承接当前运动与换道状态，由响应控制器接管；将背景动作与被测车辆动作统一送入车辆动力学。',
'S104，更新参数组选择。保持事件、策略和非选中输入一致，比较单组变更前后的有效执行，利用实际风险变化及交互相关性确定后续参数组选择概率，并保留独立条件抽样通道。',
'S105，安排下一层测试。根据有效执行结果提高风险门槛，在准入集合内联合选择交互类别代表和限制跨类别同源名额，按分配结果生成并执行后续候选。',
'S106，保存失效与复现材料。对有效执行中的实际失效保存完整输入、实际轨迹和模型版本；达到预算或终止条件后输出失效档案、测试记录及覆盖统计。可将所述档案进一步用于已有策略的定向更新。']:
    para(D,x)
heading(D,'5 技术方案的具体实施方式')
heading(D,'5.1 数据对象、固定条件与重放约定',2)
para(D,'为使行为变化能够逐组核对，本方法将固定条件、可变输入和执行结果分开保存。表1给出贯穿各步骤的数据对象；同一符号在全文中具有相同含义。')
table(D,['对象','符号','内容与用途'],[
('自然事件','$e$','道路几何、稳定车辆身份、尺寸、起始状态、有效历史及记录轨迹。'),
('初始条件／占位／尺寸','$C_0,M,D$','初始运动关系、有效背景位置掩码及全部车辆长宽；局部变更时固定。'),
('行为输入／参数组集合',r'$\omega,\mathcal B_e$','允许变化的行为输入，以及事件内可替换参数组的索引集合。'),
('环境参数／当前策略',r'$\theta,\pi$','冻结的生成、响应、控制与动力学配置，以及本轮被测策略。'),
('执行轨迹／风险／失效标签',r'$\tau,R,f$','实际联合状态序列、非负事件风险，以及按预定规则得到的二值失效标签。'),
('交互类别／独立来源',r'$g,\alpha$','依据实际交互描述得到的类别，以及独立条件抽样的谱系起点。'),
('层索引／风险门槛',r'$\ell,\gamma_\ell$','当前风险层及该层要求的最小分层得分。'),
('准入集合／实际名额总数',r'$\mathcal H_\ell,N$','满足有效性和本层门槛的候选集合，以及实际可分配的后续测试名额。')],[3.5,2.5,10.2])
para(D,'每组行为输入同时记录所属车辆、变量类型、作用时间范围及允许分布。按车辆身份、物理帧号和变量类型预先索引随机数组；某分支未使用的数组元素仍保留在原位置，避免控制分支改变随机数的消费顺序。策略存在随机动作时，另行保存并复用其随机实现。')
eq(D,r'\tau=S_{\theta}(e,\omega;\pi)',1,r'式（1）中，$S_{\theta}$为在环境参数$\theta$下执行交通仿真的映射；输入为事件$e$、行为输入$\omega$及当前策略$\pi$，输出$\tau$为实际轨迹。策略随机实现作为保存的固定执行条件，不混入待变更的背景行为参数。')
para(D,'主实施方式从事件起点重新执行原输入与变更输入，并重置控制器记忆、换道状态及参考进度。采用执行前缀复用时，须一并恢复车辆状态、历史掩码、控制器和策略内部状态、已开始的换道及全部随机索引；已影响所复用前缀的参数不得在该前缀之后直接替换。')
para(D,'失效判据在本轮开始前固定。一种实现将被测车辆碰撞和驶出道路分别记录，并以二者的逻辑或形成安全失效标签；停滞、进度和背景车辆相互碰撞另外记录。非法初始重叠、缺失关键输入及数值发散记为执行异常，保留其成本和原因，不以其代替被测策略失效。')
heading(D,'5.2 自然驾驶事件及风险尾部选择',2)
para(D,'读取车辆轨迹、车辆属性及道路元数据，将包围框位置转换为车辆中心位置，统一行驶方向和道路坐标。按照事件起始时刻确定同车道及相邻车道前后车辆的稳定身份，并截取连续窗口。车辆在后续执行中变换车道时，其前车、后车或切入者角色可以改变，身份和数组索引保持不变。')
para(D,'检查采样连续性、车辆尺寸、位置跳变、加速度、初始占用和道路有效性。换道事件保存越线前后的有效上下文。对包含共有车辆及重叠时间段的自然片段，保存自然记录的关联组标识，以便后续分组使用；该标识与独立条件抽样的来源标识分别存储。')
para(D,'事件风险由车辆几何和相对运动计算。一种实现根据车辆半长和、半宽和、纵向期望时距、接近速度及制动能力构造动态安全椭圆；在短预测窗口内，沿当前相对运动外推位置，将进入该椭圆的程度作为危险量，并对参与车辆和预测时刻进行聚合，得到每个物理时刻的非负瞬时风险。')
eq(D,r'R(\tau)=\max_{0\leq t\leq T}\Phi_t+\lambda_{\mathrm{exp}}\Delta t\sum_{t=0}^{T}\Phi_t',2,r'式（2）中，$t$为物理帧索引，$T$为最后一帧索引，$\Phi_t$为瞬时安全包络侵入风险，$\Delta t$为物理步长；非负参数$\lambda_{\mathrm{exp}}$控制持续暴露的权重。输出$R(\tau)$同时反映风险峰值和持续时间，并用于事件筛选和后续参数比较。若瞬时风险取无量纲量，则暴露权重采用每秒为单位。')
para(D,'对自然事件风险超过拟合阈值的部分拟合广义帕累托分布，并按预先确定的尾部水平选取目标事件。尾部拟合阈值、事件选择阈值和后续搜索门槛分别记录，避免将三者混用。除长尾事件外，可按事先规定的比例保留一般事件作为独立探索的事件来源。')
heading(D,'5.3 固定事件条件下的参考运动生成',2)
para(D,'选定事件后，将道路、车辆身份、车辆尺寸、起始状态及占位掩码作为条件输入；记录轨迹用于提供事件依据和重建参照，不强制背景车辆在后续每个时刻经过记录位置。行为参数分为四类：参考运动生成输入、持续驾驶员特征、换道特征，以及按车辆和时刻索引的响应随机输入。')
para(D,'在一种实现中，条件归一化流根据初始条件和占位掩码生成各背景车辆的稀疏二维位置结点。以事件起点和这些结点进行保形分段三次Hermite插值，形成连续基准位置轨迹；条件扩散模型生成联合位置残差，残差经过端点及结点约束后与基准轨迹相加，得到背景车辆的软参考运动。车辆尺寸随事件保持，不作为扩散轨迹残差的变化输入。')
eq(D,r'K=F_{\theta}(C_0,M,z_K),\quad P=P_{\mathrm{H}}(C_0,K)+E_{\theta}(C_0,M,K,z_E)',3,r'式（3）中，$K$为稀疏位置结点，$F_{\theta}$为条件结点生成映射，$z_K$为其标准基空间输入；$P_{\mathrm{H}}$为插值基准轨迹，$E_{\theta}$为条件残差生成映射，$z_E$为扩散初始噪声，$P$为输出软参考。初始条件$C_0$和占位掩码$M$在该事件的局部比较中保持不变。')
para(D,'参考运动经地图、坐标和动力学接口转换为跟踪参考，反馈控制器根据当前状态、参考速度及参考曲率计算基准动作。所述基准动作在当前实际状态上重新计算，因此保持某车辆的输入并不要求其执行轨迹与另一比较运行完全相同。联合生成模型的条件耦合和真实交互均可使其他车辆的后续状态发生变化。')
para(D,r'条件抽样律记为$\mu_e$，表示在固定事件$e$下完整生成允许行为输入的规则。独立抽样重新生成完整行为输入并赋予新的来源标识；单组局部变更继承原来源。独立抽样需要切换到其他自然事件时，另行保存所选事件标识，并重新开始该事件内的条件比较。')
heading(D,'5.4 参考适用性判断与背景车辆响应',2)
para(D,'为保持生成参考与实际交通之间的联系，每个物理时间步使用当前联合状态识别前车、侧向进入车辆、目标车道前后车辆及已经产生响应的主体。响应模型使用已经实现的状态、有效历史和本车参考；对被测车辆未来运动的预测由当前可观测状态外推获得，不读取其未来控制指令。')
para(D,'一种参考适用性判断结合短时冲突预测、经验预测误差轨迹及跟驰或前车离开后的交互刺激。在经验误差轨迹的多种实现下计算车辆几何冲突支持比例；支持比例达到预设阈值，或观测交互刺激达到相应阈值时，判定需要响应。门槛在自然数据标定和开发验证后固定，并记录触发时刻、触发主体及触发原因。')
eq(D,r'\chi_t^i=\frac{1}{J}\sum_{j=1}^{J}\mathbf{1}\{\exists h\in\mathcal T:\mathcal C_{t,h}^{i,j}=1\}',4,r'式（4）中，$i$为背景车辆索引，$J$为经验误差轨迹的实现数，$j$为实现索引，$\mathcal T$为预测时间点集合；$\mathcal C_{t,h}^{i,j}$为该实现下预测到时间偏移$h$时的几何冲突标记，$\mathbf{1}$为示性函数。输出$\chi_t^i$为触发参考适用性判断的冲突支持比例，表示经验预测支持程度。')
para(D,'纵向响应可采用按自然驾驶数据标定的智能驾驶员模型，并通过观测历史、切入关系及持续残差修正响应；驾驶员期望速度、最小间距、期望时距、加速度及舒适减速度由联合参数分布取得，单次实例内保持其持续特征。横向响应将换道决策与连续执行分开，根据目标车道实际间距和预期加速度收益选择动作，再由连续轨迹及反馈转向实现。')
para(D,'主实施方式采用单次事件内的一次持续接管。未接管车辆继续由参考跟踪器控制；接管时保存原车道、目标车道、换道是否已经开始、当前进度和上一实际动作，随后由背景驾驶策略依据实际交通控制该车至本事件结束。参考到期也可触发接管。各车辆在同一时刻的状态上完成动作判断后，统一推进动力学。')
eq(D,r'o_t^i=o_{t-1}^i\lor v_t^i,\quad a_t^i=(1-o_t^i)a_t^{i,\mathrm{base}}+o_t^i a_t^{i,\mathrm{resp}}',5,r'式（5）中，$o_t^i$为是否已经接管的二值状态，初值为0；$v_t^i$为本时刻的接管触发标记，$\lor$为逻辑或；$a_t^{i,\mathrm{base}}$为参考跟踪动作，$a_t^{i,\mathrm{resp}}$为依据实际状态生成的响应动作。两类动作先满足各自接口的可执行限制，再交由统一动力学接口执行。未接管时输出基准动作，接管状态在事件内持续保留。')
para(D,'接管标记与实际运动改变分别记录。车辆被标记为已接管，但其执行动作尚未偏离基准时，不以标记本身作为向其他车辆传播响应的证据；后续冲突预测使用实际已经实现的动作及运动。图2据此区分控制权状态、实际响应状态和交通状态更新。')
heading(D,'5.5 响应后的换道承接与参考进度处理',2)
para(D,'背景车辆因跟驰或让行减速后，原时间参考可能已经超前。主实施方式在接管后改由背景驾驶策略控制，停止以原时间参考的纵向位置要求车辆追赶；尚未开始的换道依据实际目标车道间隙重新决策，已经开始的换道从当前位置、横向速度及既有目标车道继续构造可执行轨迹，避免将车辆重新拉回换道起点。')
para(D,'在需要继续利用原空间路径的另一实施方式中，将路径形状与到达时间分开表示。实际进度接近参考时按时间查询；因交互产生明显进度偏移后，在纵向单调的路径区段按实际纵向位置重新查询横向参考，或在一般道路上按弧长投影查询。')
eq(D,r'\eta_t^i=\operatorname{clip}\left(\frac{x_t^i-\bar x_j^i}{\bar x_{j+1}^i-\bar x_j^i},0,1\right),\quad y_t^{i,\mathrm{ref}}=\bar y_j^i+\eta_t^i(\bar y_{j+1}^i-\bar y_j^i)',6,r'式（6）中，$x_t^i$为车辆实际纵向位置；$j$为包围该位置的参考路径区间索引，$\bar x_j^i,\bar y_j^i$为该路径结点坐标；$\eta_t^i$为区间内的插值比例，$y_t^{i,\mathrm{ref}}$为输出横向参考，$\operatorname{clip}$表示区间截断。仅在分母大于预设正阈值时使用该式；重复纵向位置先合并或跳过，停车时保持当前路径参考。')
para(D,'采用前述空间查询时，切换阈值、参考进度及当前换道状态一并保存，动作变化受加速度和变化率约束。是否在状态恢复相容后释放响应控制，可在事件开始前选定；该变体与一次持续接管分别配置，配对比较始终使用同一种执行方式。若追加尾段用于完成已经开始的换道，应保存尾段长度和计分规则，并在全部比较运行中一致使用。')
heading(D,'5.6 基于实际风险变化的参数组选择',2)
para(D,'实际执行给出风险最大的相关车辆、关键时刻及交互阶段。根据这些信息限定可选择的车辆和响应时间范围，并为当前事件的参数组设置交互相关权重。权重可以采用二值掩码：属于关键车辆或其直接交互车辆，且覆盖相关作用时间的组取1；不适用的组从候选集合中移除。')
para(D,'对选中的参数组执行单组变更。事件、当前策略、环境参数及其他组的数组和值保持一致；由新输入从事件起点重新执行，并允许其他车辆因实际交互产生不同响应。仅将原输入和变更输入均有效的配对结果用于更新该组敏感度；比较可复用已经执行并保存的原输入结果。')
eq(D,r'd_b=\frac{1}{L_b}\sum_{r=1}^{L_b}\min\left(\frac{|R_{b,r}^{\prime}-R_{b,r}|}{1+|R_{b,r}|},d_{\max}\right)',7,r'式（7）中，$b$为参数组索引，$L_b$为该组的有效配对数；$r$为配对索引，$R_{b,r}$和$R_{b,r}^{\prime}$分别为原输入和单组变更输入的实际事件风险，$d_{\max}$为预设截断上限。输出$d_b$衡量局部变更引起风险变化的幅度；使用绝对值是为识别具有调整影响的组，具体风险升降由后续执行和准入判据决定。')
eq(D,r's_b\leftarrow(1-\rho)s_b+\rho d_b,\quad q_b=\frac{a_b s_b+\delta}{\sum_{c\in\mathcal B_e}(a_c s_c+\delta)}',8,r'式（8）中，$s_b$为累计敏感度，$\rho$为介于0与1之间的更新权重，$a_b$为交互相关权重；$\delta$为正平滑常数，$c$为参数组求和索引。输出$q_b$为后续局部提议选择参数组$b$的概率。尚无有效配对时保留预设初始敏感度，使适用但未充分探索的组仍有机会被选择。')
para(D,'对于标准高斯基空间中的连续输入，采用预条件Crank–Nicolson型局部提议，仅更新选定车辆和时间范围内的元素；其余元素逐项保留。离散意图则按照事件允许的条件分布重采样。')
eq(D,r'\omega_b^{\prime}=\sqrt{1-\beta_b^2}\,\omega_b+\beta_b\xi_b,\quad \xi_b\sim\mathcal N(0,I),\quad 0<\beta_b\leq1',9,r'式（9）中，$\omega_b$和$\omega_b^{\prime}$为变更前后的标准基空间参数，$\beta_b$为提议尺度，$\xi_b$为同形状标准正态创新，$I$为单位协方差矩阵。该式用于连续参数的局部候选生成；映射到驾驶参数后仍检查其允许范围。')
eq(D,r'K_e(\omega,\cdot)=(1-\varepsilon)\sum_{b\in\mathcal B_e}q_bK_{e,b}(\omega,\cdot)+\varepsilon\mu_e(\cdot)',10,r'式（10）中，$K_{e,b}$为仅改变参数组$b$的提议规则，$K_e$为完整候选提议规则，$\mu_e$为第5.3节的完整条件抽样律，$\varepsilon$为严格介于0和1之间的独立抽样比例。独立抽样建立新来源，局部提议继承原来源。')
para(D,'上述规则用于定向发现失效。参数组概率随已观察风险变化而更新，候选又受风险与覆盖条件筛选，因此搜索命中比例按发现效率报告；不将该比例直接解释为自然事件分布中的失效概率。所有试运行、被拒绝提议及执行异常均记录其实际消耗。')
heading(D,'5.7 风险分层及交互描述',2)
para(D,'搜索从独立条件抽样得到的初始群体开始。将实际事件风险转换为有界分层得分，并使已经发生预定失效的有效实例进入最高层，以区分风险引导和实际失效判定。')
eq(D,r'Z_n=f_n+(1-f_n)\frac{R_n}{1+R_n},\quad \gamma_\ell=\max\left\{\gamma_{\ell-1},Q_{1-p_0}(\{Z_n\})\right\}',11,r'式（11）中，$n$为本层有效候选索引，$R_n$为事件风险，$f_n$为二值失效标签，$Z_n$为分层得分；$Q_{1-p_0}$为经验分位数，$p_0$为目标保留比例。初始门槛取0。只有满足事件约束、执行有效且得分不低于$\gamma_\ell$的候选进入准入集合$\mathcal H_\ell$；有效候选不足时按剩余预算补充。')
para(D,'从实际轨迹提取关键车辆与被测车辆的相对纵横向位置、相对速度、关键时刻、车辆角色和控制阶段。连续量按预先固定的尺度归一化；以角色和控制阶段形成上层类别，再按预设区间或网格细分。也可使用在本轮开始前固定的距离聚类。相同规则用于本轮全部候选，派生实例执行后重新计算类别。')
para(D,'一种与现有执行接口相容的描述使用关键时刻附近的五个连续量：相对纵向位置、相对横向位置、相对纵向速度、相对横向速度和关键时刻占窗口长度的比例。上层角色包括前方、后方、邻侧及自由道路，并记录跟踪或响应阶段；对驶出道路另记方向。跟驰制动、切入冲突及目标车道后车冲突可作为经过明确事件判定后的解释性类别。')
heading(D,'5.8 交互类别与独立来源的联合名额分配',2)
para(D,'风险准入只表明候选可以继续向危险条件扩展，后续还需决定哪些候选获得多少测试机会。为此，对本层每个准入候选设置非负整数名额，名额用于保留该候选或从其输入生成一次后续提议。分配以保留更多交互类别为目标，同时累计同一来源在全部类别中的名额。')
eq(D,r'\max_{\{m_n\}}\sum_{g\in\mathcal G_\ell}\mathbf1\left\{\sum_{n\in\mathcal H_\ell:g_n=g}m_n>0\right\}',12,r'式（12）中，$m_n$为候选$n$获得的名额，$g_n$为其当前交互类别，$\mathcal G_\ell$为本层准入集合中出现的交互类别集合。目标函数统计至少获得一个名额的类别数量。')
eq(D,r'\sum_{n\in\mathcal H_\ell}m_n=N,\quad m_n\in\mathbb N_0,\quad M_\alpha=\sum_{n\in\mathcal H_\ell:\alpha_n=\alpha}m_n\leq c_\alpha',13,r'式（13）中，$N$为实际分配总名额，$\mathbb N_0$为非负整数集合；$\alpha_n$为候选$n$的独立来源，$M_\alpha$为来源$\alpha$跨全部交互类别获得的名额合计，$c_\alpha$为其预设上限。同一来源进入多个类别后仍使用同一累计量。')
para(D,'一种可实施的分配过程分为代表保留和剩余填充。先建立“交互类别—独立来源”二部图：存在某类别、某来源的准入候选时连接二者。设置源点到各类别的容量为1、类别到来源的容量为1、来源到汇点的容量为其名额上限，计算最大流，并将代表保留数限制在总名额以内。每个匹配边选取一个对应候选作为类别代表；候选并列时，优先选择与已保留行为描述距离较大者，再比较风险。')
para(D,'代表保留后，在各来源剩余容量内继续填充名额，优先选择覆盖不足、描述重复较少且风险较高的候选。同一父候选可获得多个后续提议名额，各提议使用独立创新产生不同候选。类别数超过总名额时，先最大化可保留类别数，再按历史覆盖不足程度及风险排序；可通过整数规划实现相同约束。')
para(D,'来源标识仅在完整独立条件抽样时新建，样本保留、局部变更和跨类别移动均继承该标识。若准入来源的剩余总容量不足以填满计划名额，则在剩余预算内补充独立抽样，并按同一门槛检查；仍不足时缩减本层实际名额总数并记录，不重编号来源，也不通过放宽风险门槛填数。')
eq(D,r'A_{\mathrm{eff}}=\frac{N^2}{\sum_\alpha M_\alpha^2}\geq\frac{N}{c}',14,r'式（14）中，$A_{\mathrm{eff}}$为按名额占比计算的有效来源数；该下界适用于全部来源上限均不大于正整数$c$且$N>0$的情况。由于$M_\alpha^2\leq cM_\alpha$且各来源名额之和为$N$，可得到该界。它约束一次分配的来源集中程度，类别覆盖仍由式（12）独立确定。')
para(D,'按分配结果使用第5.6节提议规则产生后续候选，经同一执行模块计算风险，并按当前门槛接纳或拒绝。被拒绝的提议已经消耗一次实际测试，其失效保存不受接纳结果影响。图4说明交互类别与来源上限在同一次名额分配中的作用。')
heading(D,'5.9 失效保存、停止条件与策略更新接口',2)
para(D,'每次执行完成后，将完整实例和执行记录写入测试存储；有效执行发生失效时建立失效索引。记录包括自然事件标识、道路几何、稳定车辆身份、尺寸、起始状态、历史有效掩码、全部行为输入、生成参考或其重建输入、逐时随机数组、策略随机实现、模型及配置版本、实际轨迹、风险过程、交互类别、父实例和独立来源。内容哈希用于核对实例一致性。')
para(D,'完整测试记录与代表档案分开管理。完整记录保存实际发现的失效及复现输入；代表档案根据行为距离、容量及来源上限选取用于查看或训练的代表。代表被替换、未获得下一轮名额或被移出有限容量档案时，其已保存的失效记录继续保留。重复执行计入成本，但不增加按实例哈希去重的失效数。')
para(D,'在搜索开始前设置总执行预算、每层最大提议数、最大层数及无新增交互的停止次数。初始化、配对比较、独立补充、拒绝提议和异常运行均计入实际执行账本；同时记录各次物理步数。达到预算、最大层数或连续规定次数没有新增交互时结束搜索。若采用固定成本比较，则局部路径终止后按事先规则从新独立来源重启，直至耗尽共同预算。')
para(D,'定向更新为失效档案的可选用途。强化学习策略可从常规、新增失效及历史场景中混合重置，并由当前策略重新采集训练转移；监督规划器须使用经完整闭环执行验证的修正示范或符合其原生目标的材料。失效执行中的动作不直接作为正确动作标签。更新后的策略需要重新执行风险比较和安全检查，旧策略轨迹不能替代新策略运行结果。')
para(D,'本发明的测试输出不依赖特定更新算法，也不以跨版本成功—失败邻接对或区域转移调度作为必要步骤。系统可独立向外部开发流程输出可重放场景。')
heading(D,'6 关键技术特征与欲保护内容')
heading(D,'6.1 核心技术关系',2)
para(D,'核心保护内容为以下相互衔接的处理关系：固定自然驾驶事件的道路、车辆身份、尺寸和起始状态；以具有车辆及时间索引的行为参数组生成参考运动；在实际交互使参考不再适用时承接车辆当前运动和已开始的换道，依据实际状态执行背景响应；在同一事件、策略及非选中输入下利用单组变更前后的实际风险变化选择后续参数组；在风险准入集合内联合保留交互类别代表，并跨类别累计和限制独立来源获得的后续名额；保存有效失效的完整输入及执行结果。')
para(D,'其中，行为参数组同时用于采样、配对比较、局部变更和重放；实际交互结果同时用于确定响应对象、估计局部风险影响和形成类别描述；交互类别与独立来源约束作用于同一次测试安排。上述输入输出依赖是本技术方案的组织重点。')
heading(D,'6.2 可进一步限定的技术特征',2)
table(D,['可限定内容','对应技术特征及用途'],[
('事件条件和外生输入','固定初态、占位和尺寸；保存逐车逐时随机数组；未选中元素保持索引和值。'),
('参考适用性与一次接管','结合经验误差预测和实际交互刺激触发接管；区分接管标记与已实现运动变化；事件内持续保持控制权。'),
('响应连续性','承接既有换道目标和实际进度；或者在规定条件下按空间位置重查同一路径。'),
('实际影响引导','用有效配对的归一化风险差更新敏感度；正平滑项与独立抽样保留探索机会。'),
('联合名额约束','类别—来源二部图代表选择、剩余容量填充、跨类别来源累计及容量不足处理。'),
('记录和复现','内容哈希核验、父子谱系、全量失效记录与容量受限代表档案分离。')],[4.2,12])
heading(D,'6.3 系统、设备及技术边界',2)
para(D,'对应系统包括事件处理模块、条件行为生成模块、交互执行模块、风险比较模块、联合名额分配模块和失效存储模块。事件处理模块输出固定条件；生成模块输出参考及行为输入；交互执行模块输出实际状态与响应记录；风险比较模块输出参数组选择概率及风险准入结果；名额分配模块输出后续提议安排；存储模块保存输入、轨迹与谱系。可选策略更新模块读取失效档案，并向下一轮测试输出新策略版本。')
para(D,'各模块可在同一仿真服务器运行，也可由具有事件标识、实例哈希及模型版本的接口连接。相应计算机设备和存储程序的计算机可读存储介质可实现上述方法。图1及图2采用处理步骤和数据流表示模块关系，避免将具体模型名称或计算硬件型号作为必要限制。')
heading(D,'7 有益效果')
para(D,'事件固定条件与显式参数组共同限定每次直接变化的输入，使失效实例能够对应至原始道路、车辆关系及具体行为变更。逐车逐时随机索引和完整实例保存使配对比较与复现使用相同执行条件。')
para(D,'参考适用性判断把背景车辆的控制从记录或生成参考衔接至实际交通响应。承接已经开始的换道和当前动作，能够避免重新初始化机动造成的运动断续；接管后停止追赶原时间位置，或按空间进度重查路径，使跟驰让行与横向路径执行具有一致的实际进度依据。')
para(D,'同条件比较得到的实际风险差反映参数变更经过背景响应和被测策略后的局部影响，并直接改变下一次参数组选择。其作用是使测试投入对应于当前策略和局部交通条件中实际发生影响的行为变化。')
para(D,'风险准入将搜索引向更危险的条件；类别代表保留与跨类别同源上限分别约束交互重复和来源集中，且共同决定后续测试安排。有效来源数下界给出一次名额分配的可核对数量性质，完整失效记录则使代表选择不影响既有失效的复现。上述效果分别由事件约束、控制衔接、风险比较和联合分配机制产生。')
heading(D,'8 具体实施例与已有技术证据')
heading(D,'8.1 highD自然驾驶事件及条件行为生成',2)
para(D,'以highD高速公路自然驾驶轨迹为输入。[9] 处理器按录像号、车辆身份及起始帧关联轨迹、车辆属性与车道元数据，将车辆中心状态转换为统一道路坐标。保留25 Hz采样，按150帧间隔选取候选窗口，每窗包含150个状态和149次转移，首末状态间隔5.96 s；起始被测车辆为速度不低于5 m/s的小汽车。')
para(D,'在窗口起点，从同车道及左右相邻车道分别选取150 m内最近的前车和后车，最多设置六个背景位置。所选车辆全窗身份稳定，缺席位置以掩码表示；换道片段要求越线前至少25帧、越线及之后至少50帧上下文，并检查前后稳定片段和相邻车道关系。车辆进入仿真后的交互角色依据实际状态更新。')
para(D,'采用1.5 s短时风险预测窗，按0.2 s间隔采样并包含末端，计算安全椭圆侵入风险。式（2）的物理步长取0.04 s，持续暴露权重取每秒0.15；相同计算接口用于自然事件和实际执行轨迹。数据处理中形成96,055个合格片段，现有长尾条件模型使用2,964个目标尾部事件。该数量说明现有资产的数据规模，不作为本方法的必要数量限制。')
para(D,'对选定事件，保存真实道路几何、车辆尺寸及初态。条件向量包含26维初始条件和六位占位掩码；26维由被测车辆的纵横向速度，以及六个背景位置各自相对被测车辆的纵横向位置、速度组成。每辆背景车在2.0、4.0和5.96 s设置二维结点，共36个结点坐标量；以起点和上述结点构造插值参考，再加入条件扩散残差。')
para(D,'本事件的测试仅改变结点基输入、扩散噪声、驾驶员特征、换道特征及响应随机数组，初始条件、占位及尺寸作为条件传入。当前工程提供固定初态和占位的条件采样接口，以及上述行为输入的存取和局部变更接口；本实施例按照第5节将其连接至固定事件的测试与分配流程。')
heading(D,'8.2 一次持续接管的交互执行',2)
para(D,'将事件道路、尺寸、起始状态和参考载入交通仿真器，以0.04 s物理步长更新世界状态；被测策略按其部署频率输出动作，低于物理频率时在相应子步内保持或按规定转换动作。背景跟踪和响应接口每个物理步检查状态，保证不同策略决策频率下使用同一物理执行口径。')
para(D,'一种配置使用1.5 s参考适用性预测窗，冲突支持比例门槛为1/16，并结合已标定的跟驰和前车离开刺激。响应接管时从实际位置承接换道，记录触发车辆、时刻、原因及实际动作是否改变。上述门槛是可配置实施参数，并非每个事件必须触发接管的要求。')
para(D,'现有工程已有自然记录重建、参考适用性及人工控制干预材料：自然Test报告覆盖10,151个场景；控制干预材料包含51个预选案例、204个执行回合。既有CPU计时报告在最多七个参与实体的8,493个物理步中记录到世界更新耗时99%分位约8.59 ms，未出现超过40 ms的步骤，该计时不含被测自动驾驶策略计算。以上材料支撑执行组件及接口的可运行性；本节不据此给出联合名额分配带来的失效发现增益。')
heading(D,'8.3 参数比较与测试名额分配示例',2)
para(D,'选取关键背景车辆的一组驾驶员参数作为变更对象，固定事件、策略随机实现、参考生成输入及其他车辆参数。对原输入与变更输入分别从起点执行；有效时按式（7）计算风险变化，并按式（8）更新选择概率。连续参数局部提议尺度可取0.2，敏感度更新权重可取0.2，正平滑常数可取0.05；这些数值提供一种配置示例，实际使用前与参数分组规则共同固定。')
para(D,'现考虑一轮计划分配12个后续名额，当前准入候选具有三种交互类别，来自四次独立条件抽样A、B、C、D，每个来源上限为3。表2给出满足约束的一种分配，其中同一来源可以在多个类别中出现。该表为分配规则的算例，不作为仿真实测结果。')
table(D,['交互类别','来源A','来源B','来源C','来源D','名额合计'],[
('跟驰制动','2','2','0','0','4'),('切入冲突','1','0','2','1','4'),('目标车道后车冲突','0','1','1','2','4'),('来源合计','3','3','3','3','12')],[5.2,2,2,2,2,3])
para(D,'三个类别均获得代表，各来源跨类别合计均为3，有效来源数为4。若准入来源减少为三个且上限仍为3，则最多可分配9个名额；补充抽样经同一风险门槛接纳后方可增加来源容量，否则本层实际名额缩减为9。派生实例改变交互类别时继续沿用原来源。')
para(D,'独立条件抽样比例可取0.1，风险层目标保留比例可取0.1或0.25；一轮内固定相应配置。风险门槛、参数组影响更新、类别网格和来源上限共同记录，保证每次提议能够追溯至其父实例、所选参数组及获得名额的原因。')
heading(D,'8.4 失效复现与效果验证口径',2)
para(D,'实际失效发生后，实例加载器恢复固定事件、生成参考或其完整重建输入、驾驶员与换道参数、逐时响应随机量、策略版本及配置，从事件起点重新执行，并比较输入哈希、关键风险时刻和失效结果。用于策略更新时保持同一实例恢复接口，重新采集交互或验证修正示范。')
para(D,'对联合方法的技术效果，采用相同被测策略、相同事件池、相同执行配置和相同实际预算，比较随机条件抽样、风险分层搜索及本方法。主要观察按实例哈希去重的失效数量、失效所对应的独立来源数量、交互类别覆盖、名额集中程度和实际执行成本；分别移除参数组影响更新或同源约束，以判断相应机制的作用。该验证口径为后续完整实施的核验方案。')
para(D,'已完成的组件证据、完整联合设计及当前代码接通状态在随附《聚焦检索与工程实现核对说明》中逐项对应。本次修订使用已有记录，不新增搜索、训练或性能测量，也不将历史不同预算或不同策略结果合并为本发明的对比结论。')
heading(D,'9 替代实施方式及应用')
para(D,'事件筛选可由极值尾部拟合替换为预设分位数或明确的关键事件条件；条件行为生成可由其他保持相同固定条件和参数组接口的概率模型实现。参考适用性判断可采用几何预测、经验误差支持或经标定的交互刺激，响应控制可使用参数化驾驶模型、学习策略或二者组合。')
para(D,'交互类别可由固定网格、确定性聚类或交通类型与连续描述共同形成。来源上限可统一设置，也可按预先规定的规则分来源设置；无论采用何种上限，局部变更须继承谱系，并在同一次分配中跨类别累计。联合分配可以采用最大流、整数规划或满足相同约束的确定性过程。')
para(D,'可选扩展至城市道路或其他交通参与者时，应补充对应的道路拓扑、可执行运动模型、事件筛选及交互判据，并保持固定条件、参数替换、实际执行比较和联合分配之间的关系。该扩展属于可选替代方案，待发明人结合实际实现确认。')
heading(D,'10 附图占位及绘制说明')
para(D,'附图采用黑白或灰度线框，以方法步骤、判断条件和输入输出为主要内容。现阶段保留图1—图4的占位与绘制说明，后续绘图沿用正文编号。')
for title,desc in [
('图1 自然驾驶事件约束的多失效域测试流程','从左至右或自上而下绘制S101事件构建、S102条件行为生成、S103交互执行、S104参数组选择、S105联合名额分配、S106失效保存。事件固定条件分别连接生成和执行模块；实际风险反馈连接参数组选择与风险准入；来源标识连接分配和存储。独立抽样支路取得新来源，局部变更支路继承来源。策略更新以虚线可选支路表示。'),
('图2 参考适用性判断、一次接管与机动承接','绘制参考跟踪器、联合实际状态、参考适用性判断、持续接管状态、背景响应策略及统一动力学。未接管支路输出基准动作；触发后继承当前换道目标、进度及上一动作。单独画出“控制权已接管”和“实际动作已经改变”的记录节点，实际运动反馈至下一步关系判断。空间路径查询作为可选分支加注。'),
('图3 同条件参数比较与后续提议','并列绘制原输入运行与仅替换一组的运行，两者共享事件、策略、环境配置及其余随机数组。两条实际轨迹进入风险差计算，继而更新参数组敏感度及选择概率。标注车辆掩码、时间范围、连续局部提议和独立完整抽样；最终候选返回图1的同一执行模块。'),
('图4 交互类别与独立来源的联合名额分配','上部绘制风险准入集合及其类别—来源二部图，类别侧标记代表容量1，来源侧标记名额上限。下部以表2的12格名额示意跨类别累计，同一来源统一用字母A、B、C、D及纹理识别。右侧给出来源不足时补充抽样或缩减实际名额的分支；失效存储单独连接，表示不受后续名额筛选影响。')]:
    heading(D,title,2)
    p=para(D,'【附图占位：待绘制】');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
    para(D,desc)
D.add_page_break()
heading(D,'主要公开参考文献')
references(D)
MAIN=OUT/'技术交底书_自然驾驶事件约束的自动驾驶多失效域测试方法_修订稿.docx'
D.save(MAIN)

# Independent, review-facing evidence note. No review logs are embedded in the invention narrative.
A=setup('聚焦检索与工程实现核对说明')
p=A.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;rich(p,'聚焦检索与工程实现核对说明',18,True)
para(A,'对应文件：技术交底书_自然驾驶事件约束的自动驾驶多失效域测试方法_修订稿.docx。核对日期：2026年10月23日。本说明用于发明人及专利代理师审阅，集中记录发明叙事、公开方案和工程支持之间的对应，不替代交底书正文。')
heading(A,'1 修订采用的技术主线')
para(A,'在保留原稿“固定自然事件＋行为参数受控变更＋交互执行＋类别与来源联合分配”构思的基础上，将创新表述集中到步骤之间的功能依赖。生成模型、驾驶员模型、选择性响应、危险搜索和多样性档案均不单独作为排他性创新主张。相应修改是对原有构思的可实施化表达，不把未连接的接口认定为已经完成的系统。')
para(A,'当前工作区的最新方法设计文档将研究重点调整为风险分层测试与策略反馈修复，并写明暂不主张多失效域识别或充分覆盖。由于本次任务明确指定本交底书，修订稿仍保留其原有测试发明构思，将策略修复作为可选用途，并将代码中已经实现的搜索与本交底书设计分别核对。该取向需要发明人最终确认。')
heading(A,'2 原稿主要修改与理由')
table(A,['原稿处理','修订处理与理由'],[
('把自然事件固定条件与通用生成搜索合写','区分固定事件条件采样接口和当前全先验生成搜索，主实施方式明确禁止变更初态、占位和尺寸。'),
('把交互类别、搜索分支和失效域近似混用','分别定义经验失效区域、交互类别、计算分支和独立来源，类别覆盖不直接充当失效域识别结论。'),
('响应释放、路径重索引与现有接管混写','以当前一次持续接管和已开始换道承接为主；空间重索引及控制释放保留为原稿支持的可选设计。'),
('风险敏感度停留在均值差的原则描述','加入归一化、截断和滑动更新，与现有单块风险变化更新机制对齐；严格固定事件和有效配对作为本发明约束。'),
('以单步高斯提议暗示整体分布保持','明确高斯局部提议只是候选生成规则，自适应选择及覆盖筛选后的命中率用于发现效率。'),
('名额优化缺少求解与不可行分支','补齐类别—来源二部图最大流、剩余填充、跨类别计数和容量不足处理。'),
('所有失效保存与有界档案保留混同','区分完整测试存储、失效索引与容量受限代表档案，避免代表数量充当发现总量。'),
('把修复混合采样作为普遍策略训练方式','区分强化学习重新采集与监督策略经验证示范；移除无证据的通用修复效果表述。'),
('以图像承载原有附图','改为图1—图4占位及技术绘制说明；公式全部使用原生Word数学对象。')],[5.2,11])
heading(A,'3 工程实现逐项对应')
para(A,'核对对象为当前工作区实际文件，不按某个提交标签推定；工作区已有未提交改动保持原样。本次只读取相关代码和已有报告，没有启动正式仿真、训练、下载数据或修改工程功能。状态分为“已有直接支持”“已有组件，联合流程待接通”和“原稿设计，当前未核到对应实现”。')
checks=[
('事件提取与风险','已有直接支持','process_highD/src/natural_segments.py；preprocess.py；safety_envelope_risk.py；config.yaml','25 Hz、150状态、5.96 s；初态、车辆、道路与风险；源事件关联组不同于独立抽样谱系。'),
('固定条件生成','已有直接支持，联合流程待接通','normalizing_flow/src/scenario.py：sample_constraints、sample_constraints_from_base_randomness；hierarchical_world_model/src/composition.py：sample_constraints','可固定初态、占位及尺寸生成结点；通用组合器可能施加scope条件处理，接入前需核对其是否改写原始条件，不能仅因函数名便认定事件全程固定。'),
('通用生成搜索','与本发明固定事件模式不同','ads/common/instances.py：GeneratedSource._materialise、mutation；ads/search/engine.py：_mutation','当前使用生成道路，并允许scenario_uniform与c0_base_latent变更；可改变占位、尺寸或初态，也可能缺少原自然事件身份。不能将其直接记为固定自然事件实现。'),
('显式外生输入与局部变更','已有直接支持','hierarchical_world_model/src/randomness.py：WorldExogenousState、mutate','支持固定数组、驾驶员与换道组、响应时间范围及车辆掩码；本发明固定事件模式须屏蔽场景及初态组。'),
('条件参考生成','已有直接支持','hierarchical_world_model/src/composition.py；diffusion/src/data.py；normalizing_flow/src/scenario.py','条件结点、插值参考及扩散残差可复用；需保持事件地图、尺寸、历史与生成输入的对应。'),
('参考适用性及接管','已有直接支持','plan_applicability.py：forecast_bundle、conflict_scores、interaction_scores；takeover.py：_admission、forward','存在冲突支持比例、实际交互刺激、一次持续接管；owned与response_executed分开保存。'),
('换道承接','已有直接支持','takeover.py：_update_plan_maneuver、_adopt_current_maneuver；stochastic_drivers/online.py','接管从实际状态承接已开始换道；当前主实现接管后由NPC控制，不返回原时序参考。'),
('空间参考重索引与释放','原稿设计，当前未核到相同实现','原稿第5.5节；当前tracking.py与takeover.py','式（6）及控制释放保留为可选实施，不称为当前跟踪器的既有功能。'),
('旧版事件匹配时序标定','未核到原稿所述统一实现','原稿第5.4节；当前stochastic_drivers/response_data.py、driver_policy.py、NPC标定文件','原稿0.2/0.4/0.8 s刺激—参照匹配与当前状态响应拟合不等同。正文改为现有响应标定及经验误差支持，不继续写成已标定步骤。'),
('风险变化引导选择','已有局部机制，严格条件待接通','ads/search/engine.py：_proposal_configuration、_advance','代码按单块提议的风险差更新敏感度，采用0.8旧值＋0.2新值，默认1，平滑0.05；车辆掩码及关键时间范围已具备。现有比较不保证固定自然事件，相关性通过掩码落实，非原稿通用乘积模型。'),
('独立抽样与联合变更','已有机制，含义须区分','engine.py：_fresh、_mutation、run；instances.py：prior_identity、mutation','joint_fraction=0.2为多块联合变更，继续继承原来源；fresh/restart才是新独立来源。不能把该0.2当独立抽样比例。'),
('分层和分支','已有直接支持','ads/search/engine.py：_advance；ads/search/branches.py：HillValleyPartitioner','存在分位风险门槛、中点检查及分类分支；属于NSuS技术基础，分类分支不是交互类别或已证明独立失效域。'),
('类别与来源联合名额分配','原稿设计，当前未接通','对照engine.py：_advance；archive.py：diverse_representatives、add','当前来源上限用于各domain的有限代表选择，跨domain没有统一来源计数，也未见主群体名额的类别—来源最大流。图4及式（12）—（14）是待实现的核心设计。'),
('全量失效与代表档案','已有部分直接支持，需核对持久化','instances.py：InstanceStore；ads/method_c.py：on_evaluation回调及持久化；archive.py','已有实例存储、记录及有限代表。archive.add排除branch_check且按域限容；本发明要求所有有效失效完整保存时，需另核诊断点及异常路径，不能单靠archive.json证明。'),
('策略修复','已有算法接口，非本发明必要步骤','ads/dqn/repair.py；ads/method_c.py：train；doc/method_c_experiments.md','DQN旧入口有0.5/0.3/0.2混合。最新五规划器的保持、局部更新与准入设计尚未全部接通，不以旧混合比例概括最新所有策略。')]
for name,state,path,note in checks:
    heading(A,name,2);para(A,'核对状态：'+state+'。');para(A,'代码或材料：'+path+'。');para(A,note)
heading(A,'4 聚焦检索范围与记录')
para(A,'检索日期：2026年10月23日。来源包括Google Patents公开专利文本、作者或机构论文页、arXiv全文、CVF论文及UCL机构库。专利比对直接读取公开文本和所列权利要求；本次未进行同族、法律状态和审查档案的完整核验。')
para(A,'外部检索使用领域通用词、已公开文献题名和公开号；未上传交底书、未发送源代码或实验数据，也未使用交底书全文作查询。范围限于自然驾驶事件仿真、选择性背景响应、关键车辆引导、危险场景多样性、分层搜索及来源分散，不扩展至一般感知、整车安全或无关生成任务。')
table(A,['查询组','代表检索词或入口'],[
('已知公开方案核验','WO2025090341A1；US20260134173A1；CaDRE Controllable Diverse driving scenarios；Niching Subset Simulation'),
('领域内补查','autonomous driving scenario generation diversity particle ancestry subset simulation'),
('公开专利补查','site:patents.google.com 自动驾驶 测试 场景 多样性 自然驾驶 扰动'),
('来源机制补查','site:patents.google.com driving scenario generation lineage diversity'),
('原稿引文核验','MixSim；LitSim；STRIVE；CCDiff；CounterScene；ReactSim-Bench及原稿所列公开页')],[4,12.2])
para(A,'最终正文保留10项直接相关参考：其中[3]与[10]属于关联专利与论文，不作为两项互相独立的创新佐证。原稿LitSim、CounterScene和ReactSim-Bench的公开入口已核验；SAFE-SIM等非核心背景在本次正文中不再扩写，避免背景挤占发明方案。CN120524809A已筛查，其一般条件生成内容未列入主要对比链。未出现匹配结果不等于不存在相关技术。')
heading(A,'5 最接近公开方案及区别分析')
# Keep individual prior-art summaries concise and grounded in primary text.
for title,known,diff in [
('[1] WO2025090341A1（2025-05-01）','公开文本及独立权利要求涉及日志仿真、交通条件触发背景行为切换和性能度量；说明书[0052]附近及空间路径跟随内容还涉及保持原场景及响应传播。','选择性响应、空间路径和事件保持应作为背景基础。修订稿需强调固定事件下的单组实际风险比较，以及其结果进入类别—来源联合名额分配的处理关系；不能把“需要时响应”独立表述为新颖点。'),
('[2] MixSim（CVPR 2023）','论文通过路线目标与低层行为策略构建响应交通，并优化相关路线获取危险变体。','本发明不能以“闭环反馈搜索危险路线”概括差异；应具体落到多类可替换行为组、同条件比较及谱系约束。'),
('[3][10] US20260134173A1／CCDiff','专利权利要求1和20涉及交互结构、关键主体排序及针对所选主体的扩散梯度引导；论文说明相关结构引导生成。','本发明比较的是保存输入中单组替换经实际执行产生的风险变化，并据此分配后续提议；还要求跨类别累计独立来源名额。仅用“关键车辆优先”不能区分。'),
('[4] CaDRE（2024预印本／ICRA 2025）','论文第III节给出真实轨迹有界扰动、质量—多样性网格和邻域空网格感知重启。','交互差异、行为网格和覆盖感知重启本身已有公开。修订稿的区别落在来源谱系随局部变更继承，并在风险准入后与类别代表共同约束本次测试名额；仅扩大档案类型数量不构成该区别。'),
('[5] Niching subset simulation（2025）','公开论文利用分区、支持向量机及递归子集模拟应对多峰性能函数。','分位门槛、分支及多峰搜索均不独立主张。与本发明相关的具体增量应由事件条件、背景响应、实际风险比较和联合分配共同限定。'),
('[6] CN116361175A（2023-06-30）','权利要求1包括自然驾驶数据清洗、分类、特征选择、关联细分、组合测试及风险域划分。','本发明的经验失效域不等同于依固定指标划分安全域；重点是可响应的固定事件行为变更和后续名额约束，不以“多域”名称形成区别。'),
('[7] CN121743096A（2026-03-27）','权利要求1及4涉及测试反馈、决策不确定性敏感区域和参数调整；其他权利要求涉及生成模型优化及新场景识别。','反馈后再生成危险场景已有近邻。修订稿区别是同条件单组替换的实际风险证据，及类别—来源联合约束，不能仅用“反馈闭环”或“敏感区域”主张创造性。')]:
    heading(A,title,2);para(A,'已公开内容：'+known);para(A,'本次区别分析：'+diff)
heading(A,'6 新颖性与创造性的起草判断')
para(A,'本次有限范围检索未在已核对的单一公开文本中确认如下完整组合：固定自然驾驶事件条件；按可重放参数组作同条件实际风险比较；以比较结果选择后续组级提议；在风险准入群体的同一次后续名额分配中联合保留交互类别并跨类别限制同一抽样来源；独立保存实际失效。该结论用于确定起草重心，不等同于排除所有现有技术或确认可授权。')
para(A,'创造性论证应说明这些特征为何需要协同：若只保留事件却改变多个输入，风险反馈不能对应具体行为；若只提高危险程度，搜索可能集中于少数模式；若只保留类别，单一谱系仍可跨类别占用大量名额。统一输入接口使实际风险比较可重放，而类别—来源联合约束使该比较所引导的后续测试保持可核对的分散条件。')
para(A,'可能的组合路径包括选择性响应仿真与已有质量—多样性搜索的结合，以及NSuS与谱系分散的结合。为使区别具备技术内容，独立方法需保留参数组—实际风险比较—联合名额安排之间的实际数据流；仅将若干已知算法并列、以抽样谱系替代交通多样性，或只在代表档案上限容，均不足以表达修订稿的联合设计。')
para(A,'当前尚无对应完整联合设计的等预算仿真增益结果。正文保留由步骤直接产生的可追溯、可重放和名额集中度约束，不写失效发现率、事故减少率或跨策略提升百分比。公式给出的来源数下界属于分配性质，与未观察失效域的覆盖保证分别处理。')
heading(A,'7 保护特征支持关系')
table(A,['候选保护特征','交底书支持','现有实现状态'],[
('自然事件固定条件及参数组替换','5.1—5.3；8.1；图1、图3','条件生成和数组接口已有；搜索接入待完成。'),
('参考不适用时响应并承接换道','5.4—5.5；8.2；图2','一次持续接管已有；空间重索引为可选设计。'),
('同条件实际风险影响选择后续组','5.6；8.3；图3','单块风险更新已有；固定事件比较的完整保证待接通。'),
('风险准入与类别—来源联合名额','5.7—5.8；8.3；图4','分层已有；联合名额求解与全局来源计数待实现。'),
('失效完整保存和重放','5.9；8.4；图1','已有存储、谱系和代表档案；全量保存范围须核对。')],[5.4,4.5,6.3])
heading(A,'8 发明人待确认事项与交付状态')
for s in [
'技术事实一：本次保留原稿的固定事件与联合名额设计，是否与发明人的当前保护意图一致；最新研究文档的风险修复主线是否另行交底。',
'技术事实二：是否已有未纳入当前工作区的固定事件搜索入口、跨类别名额分配实现或对应结果；如有，可补充为主实施例证据。',
'技术事实三：确认可选空间路径重索引及控制释放是否继续作为替代实施；确认最终失效判据、类别规则、来源上限和完整失效保存范围。',
'封面信息：原稿发明人、撰写人及联系方式仅保留在本地交付稿；申请人及电子邮箱待补充。',
'起草状态：已完成技术交底修订、聚焦检索、代码对应和附图说明。实际实现与验证状态按本说明区分；本次不代替系统实现或正式实验。',
'提交前工作：发明人确认核心事实后，由专利代理师结合申请日、完整现有技术及关联申请核定权利要求和申请材料。当前文件为交底起草成果，非提交就绪的完整申请包。']:
    para(A,s)
heading(A,'9 Word与版式检查')
para(A,'文档使用现代DOCX格式，所有行内及独立数学表达均生成为原生Office数学对象；解释文字置于公式外。交付前执行ZIP/XML、兼容模式、旧式公式与图片残留检查，并通过LibreOffice载入和PDF渲染核验分页。当前环境无Microsoft Word桌面程序，因此不将LibreOffice载入等同于Word原生重开验证。具体结构检查结果见support目录中的QA记录。')
A.add_page_break()
heading(A,'公开参考与核验链接')
references(A)
A.save(OUT/'聚焦检索与工程实现核对说明.docx')
# Minimal snapshot manifest so future reviewers can identify exactly what was read.
paths=[
'external_resources/patents/技术交底书_自然驾驶事件约束的自动驾驶多失效域测试方法_正式稿.docx',
'ads/common/instances.py','ads/common/evaluation.py','ads/search/engine.py','ads/search/branches.py','ads/search/archive.py',
'hierarchical_world_model/src/randomness.py','hierarchical_world_model/src/composition.py',
'hierarchical_world_model/src/takeover.py','hierarchical_world_model/src/plan_applicability.py',
'normalizing_flow/src/scenario.py','process_highD/config.yaml','process_highD/src/safety_envelope_risk.py',
'doc/method_c_experiments.md','results/hierarchical_world_model/bridge/latency.json','results/evt/event_bank/summary.json']
manifest={'date':'2026-10-23','source_files':{x:hashlib.sha256(Path(x).read_bytes()).hexdigest() for x in paths},'references':[{'id':n,'description':t,'url':u} for n,t,u in REFS],'assumption':'保留原稿核心设计，区分代码已实现内容与未接通的联合设计；未收到更改此取向的确认。'}
(OUT/'support/source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(MAIN)
