from pathlib import Path
import sys,re,json,hashlib
sys.path.insert(0,'/tmp/patent_drafting_deps')
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from lxml import etree
import latex2mathml.converter

OUT=Path(__file__).resolve().parents[1]
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def font(run,size=11,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    rpr=run._element.get_or_add_rPr();fs=rpr.rFonts
    fs.set(qn('w:eastAsia'),'Noto Serif CJK SC');fs.set(qn('w:ascii'),'Times New Roman');fs.set(qn('w:hAnsi'),'Times New Roman')

def math(tex):
    """Convert the small presentation-MathML vocabulary used here to native OMML."""
    root=etree.fromstring(latex2mathml.converter.convert(tex).encode())
    def E(tag,**attrs):
        e=etree.Element('{'+M+'}'+tag)
        for k,v in attrs.items():e.set('{'+M+'}'+k,str(v))
        return e
    def run(txt,style='i',variant=''):
        e=E('r');pr=E('rPr')
        if variant in ('script','double-struck','fraktur'):
            pr.append(E('scr',val=variant));pr.append(E('sty',val='p'))
        else:pr.append(E('sty',val='b' if variant=='bold' else style))
        e.append(pr)
        wpr=etree.SubElement(e,'{'+W+'}rPr')
        fs=etree.SubElement(wpr,'{'+W+'}rFonts')
        fs.set('{'+W+'}ascii','Cambria Math');fs.set('{'+W+'}hAnsi','Cambria Math')
        sz=etree.SubElement(wpr,'{'+W+'}sz');sz.set('{'+W+'}val','22')
        tx=E('t');tx.text=txt
        if txt.startswith(' ') or txt.endswith(' '):tx.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        e.append(tx);return e
    def container(tag,items):
        e=E(tag)
        for x in items:e.append(x)
        return e
    def conv(n,variant=''):
        tag=etree.QName(n).localname
        variant=n.get('mathvariant',variant)
        cs=list(n)
        if tag in ('math','mrow','mstyle','semantics','mpadded'):
            def delim(left,right,children):
                e=E('d');pr=E('dPr');pr.append(E('begChr',val=left));pr.append(E('endChr',val=right));e.append(pr)
                inner=[]
                for c in children:inner+=conv(c,variant)
                e.append(container('e',inner));return e
            if len(cs)>1 and etree.QName(cs[0]).localname=='mo' and etree.QName(cs[-1]).localname=='mo':
                left,right=cs[0].text,cs[-1].text
                if (left,right) in [('(',')'),('[',']'),('{','}'),('|','|')]:
                    return [delim(left,right,cs[1:-1])]
            out=[];idx=0
            while idx<len(cs):
                c=cs[idx]
                if etree.QName(c).localname=='mo' and c.text=='|':
                    end=next((j for j in range(idx+1,len(cs)) if etree.QName(cs[j]).localname=='mo' and cs[j].text=='|'),None)
                    if end is not None:
                        out.append(delim('|','|',cs[idx+1:end]));idx=end+1;continue
                out+=conv(c,variant);idx+=1
            return out
        if tag in ('mi','mn','mo','mtext','ms'):
            return [run('∼' if n.text=='~' else (n.text or ''), 'i' if tag=='mi' and variant!='normal' and n.text!='~' else 'p', variant)]
        if tag=='mspace':return [run(' ','p')]
        if tag in ('msub','msup','msubsup'):
            e=E({'msub':'sSub','msup':'sSup','msubsup':'sSubSup'}[tag]);e.append(container('e',conv(cs[0],variant)))
            if tag in ('msub','msubsup'):e.append(container('sub',conv(cs[1],variant)))
            if tag in ('msup','msubsup'):e.append(container('sup',conv(cs[-1],variant)))
            return [e]
        if tag=='mfrac':
            e=E('f');e.append(container('num',conv(cs[0],variant)));e.append(container('den',conv(cs[1],variant)));return [e]
        if tag in ('msqrt','mroot'):
            e=E('rad');pr=E('radPr');pr.append(E('degHide',val=1 if tag=='msqrt' else 0));e.append(pr)
            e.append(container('deg',[] if tag=='msqrt' else conv(cs[1],variant)))
            items=[]
            for c in (cs if tag=='msqrt' else cs[:1]):items+=conv(c,variant)
            e.append(container('e',items));return [e]
        if tag=='mover' and ''.join(cs[1].itertext()) in ('¯','‾','^','→','ˆ'):
            e=E('acc');pr=E('accPr');accent=''.join(cs[1].itertext())
            pr.append(E('chr',val={'¯':'\u0305','‾':'\u0305','^':'\u0302','ˆ':'\u0302','→':'\u20d7'}[accent]));e.append(pr)
            e.append(container('e',conv(cs[0],variant)));return [e]
        if tag in ('munder','mover','munderover'):
            base=conv(cs[0],variant)
            if tag in ('munder','munderover'):
                lo=E('limLow');lo.append(container('e',base));lo.append(container('lim',conv(cs[1],variant)));base=[lo]
            if tag in ('mover','munderover'):
                up=E('limUpp');up.append(container('e',base));up.append(container('lim',conv(cs[-1],variant)));base=[up]
            return base
        if tag=='annotation':return []
        raise ValueError('Unsupported MathML tag '+tag+' in '+tex)
    out=E('oMath')
    for x in conv(root):out.append(x)
    return out

def rich(p,text,size=11,bold=False):
    parts=re.split(r'(\$[^$]+\$)',text)
    for s in parts:
        if s.startswith('$') and s.endswith('$'):p._p.append(math(s[1:-1]))
        elif s:font(p.add_run(s),size,bold)

def para(d,text,style=None):
    p=d.add_paragraph(style=style);rich(p,text)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.35
    if not style:p.paragraph_format.first_line_indent=Cm(.74)
    return p

def heading(d,text,level=1):
    p=d.add_paragraph(style='Heading '+str(level));rich(p,text,14 if level==1 else 12,True)
    p.paragraph_format.space_before=Pt(12 if level==1 else 8);p.paragraph_format.space_after=Pt(6)
    return p

def table(d,headers,rows,widths=None):
    t=d.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid';t.autofit=False
    if widths:
        for col,width in zip(t.columns,widths):col.width=Cm(width)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i];c.text='';rich(c.paragraphs[0],h,10,True)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');c._tc.get_or_add_tcPr().append(sh)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row):
            c.text='';rich(c.paragraphs[0],str(txt),10)
    for r in t.rows:
        for i,c in enumerate(r.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:c.width=Cm(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.15
        cp=OxmlElement('w:cantSplit');r._tr.get_or_add_trPr().append(cp)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def eq(d,tex,n,explain):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    mp=OxmlElement('m:oMathPara');mp.append(math(tex));p._p.append(mp)
    font(p.add_run('  （'+str(n)+'）'),10)
    para(d,explain)

def setup(title):
    d=Document();s=d.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21)
    s.top_margin=Cm(2.3);s.bottom_margin=Cm(2.2);s.left_margin=Cm(2.4);s.right_margin=Cm(2.4)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    for n in ['Heading 1','Heading 2','Heading 3']:
        d.styles[n].font.color.rgb=RGBColor(0,0,0)
    settings=d.settings.element
    compat=settings.find(qn('w:compat'))
    if compat is None:compat=OxmlElement('w:compat');settings.append(compat)
    for el in list(compat):
        if el.get(qn('w:name'))=='compatibilityMode':compat.remove(el)
    x=OxmlElement('w:compatSetting');x.set(qn('w:name'),'compatibilityMode');x.set(qn('w:uri'),'http://schemas.microsoft.com/office/word');x.set(qn('w:val'),'15');compat.append(x)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('— '),9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);font(footer.add_run(' —'),9)
    d.core_properties.title=title;d.core_properties.author='蒋拯民';d.core_properties.subject='自然驾驶事件约束的自动驾驶多失效域测试'
    return d
