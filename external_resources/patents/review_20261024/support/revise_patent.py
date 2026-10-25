from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import json, re, sys, hashlib
sys.path.insert(0, '/tmp/patent_drafting_deps')
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as E

ROOT = Path('/home/hp/FITWMAMS')
OUT = Path(__file__).resolve().parents[1]
PREV = ROOT/'external_resources/patents/revision_20261023'
SOURCE = PREV/'技术交底书_自然驾驶事件约束的自动驾驶多失效域测试方法_修订稿.docx'
PRIOR = next(PREV.glob('XJ*.docx'))
helpers = (PREV/'support/build_disclosure.py').read_text().split("D=setup(")[0]
exec(compile(helpers, 'previous_helpers', 'exec'), globals())
OUT = Path(__file__).resolve().parents[1]
NS = {'w':W,'m':M}

REFS = [
('1','Mendoza J P, Unnikrishnan R, Wahal S, Zhang L. Selectively intelligent actors in simulations. WO2025090341A1，2025-05-01公开。','https://patents.google.com/patent/WO2025090341A1/en','权利要求1—3及涉及空间路径跟随、交互传播的说明书段落'),
('2','Suo S, Wong K, Xu J, et al. MixSim: A Hierarchical Framework for Mixed Reality Traffic Simulation. CVPR，2023。','https://waabi.ai/research/mixsim','方法中的路线条件响应策略及安全关键路线变体'),
('3','Lin H, Chen H, Wolff E. System and method of causal composition diffusion for closed loop traffic generation. US20260134173A1，2026-05-14公开。','https://patents.google.com/patent/US20260134173A1/en','权利要求1、20；中国同族线索CN122053616A另行核对'),
('4','Lin H, Huang X, Phan-Minh T, et al. Causal Composition Diffusion Model for Closed-loop Traffic Generation. CVPR，2025；arXiv:2412.17920，首版2024-12-23。','https://arxiv.org/html/2412.17920v3','第4节因果排序与组合引导；本次核对v3，2025-05-26'),
('5','Huang P, Ding W, Stoler B, et al. CaDRE: Controllable and Diverse Generation of Safety-Critical Driving Scenarios Using Real-World Trajectories. arXiv:2403.13208，首版2024-03-19；核对v2，2024-12-12。','https://arxiv.org/html/2403.13208v2','第III节质量—多样性定义、算法1及占用感知重启'),
('6','Rempe D, Philion J, Guibas L, Fidler S, Litany O. Generating Useful Accident-Prone Driving Scenarios via a Learned Traffic Prior. CVPR，2022。','https://research.nvidia.com/labs/toronto-ai/STRIVE/','潜空间事故生成、解场景优化及碰撞类型聚类'),
('7','白先旭，李亲，李维汉，石琴。一种自动驾驶车辆在不同安全域的测试场景创建方法。CN116361175A，2023-06-30公开。','https://patents.google.com/patent/CN116361175A/zh','权利要求1的特征分类、组合测试及安全域划分'),
('8','祝亚运，单文君，余福珍。基于混合专家模型的自动驾驶测试场景生成方法和系统。CN121743096A，2026-03-27公开。','https://patents.google.com/patent/CN121743096A/zh','权利要求1、10的敏感区域识别与场景参数调整'),
('9','Pathrudkar S, Bhaskaran S, Indla V S, et al. Realistic test scenario variation and test coverage generation for testing of autonomous vehicles. WO2025190506A2，2025-09-18公开。','https://patents.google.com/patent/WO2025190506A2/en','权利要求1—2及后续有关场景覆盖生成的权利要求'),
('10','Liang L, Cheng X, Chen T Y, Zheng X. From Particles to Perils: SVGD-Based Hazardous Scenario Generation for Autonomous Driving Systems Testing. PACMSE，2026；arXiv:2604.18918，2026-04-20公开。','https://arxiv.org/html/2604.18918v1','第3.2—3.3节自适应种子、危险模型及粒子排斥')]

def txt(el): return ''.join(x.text or '' for x in el.iter(qn('w:t')))
def ps(d): return d.paragraphs
def find(d, prefix):
    matches=[p for p in ps(d) if p.text.startswith(prefix)]
    assert len(matches)==1,(prefix,len(matches))
    return matches[0]
def rewrite(p, text):
    for child in list(p._p):
        if child.tag!=qn('w:pPr'):p._p.remove(child)
    rich(p,text)
    return p
def remove(el): el.getparent().remove(el)
def generated(fn):
    t=Document();fn(t)
    for h in t._element.xpath('//w:hyperlink'):
        old=h.get(qn('r:id'))
        if old:
            rel=t.part.rels[old]
            h.set(qn('r:id'),D.part.relate_to(rel.target_ref,rel.reltype,is_external=rel.is_external))
    return [deepcopy(e) for e in t._element.body if e.tag!=qn('w:sectPr')]
def replace_section(d,start,end,fn):
    a=find(d,start)._p;b=find(d,end)._p if end else d._element.body.find(qn('w:sectPr'))
    nodes=generated(fn)
    for n in nodes:a.addprevious(n)
    node=a
    while node is not b:
        nxt=node.getnext();remove(node);node=nxt
def after(p,fn):
    a=p._p
    for n in generated(fn):a.addnext(n);a=n
def figure(d,n,title,description):
    p=para(d,f'图{n}  {title}');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
    for r in p.runs:r.bold=True
    p=para(d,'【附图占位：待绘制】');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
    p=para(d,description);p.paragraph_format.space_after=Pt(9)
    return p
def link(p,label,url,size=9):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    f=OxmlElement('w:rFonts');f.set(qn('w:ascii'),'Times New Roman');f.set(qn('w:hAnsi'),'Times New Roman');f.set(qn('w:eastAsia'),'Noto Serif CJK SC');pr.append(f)
    s=OxmlElement('w:sz');s.set(qn('w:val'),str(int(size*2)));pr.append(s)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def refs(d):
    heading(d,'2.5 本节引用的相近公开文献',2)
    for n,bib,url,loc in REFS:
        p=para(d,f'[{n}] {bib}');p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(3)
        for r in p.runs:r.font.size=Pt(9.5)
        link(p,'  原文链接',url)
        p=para(d,'核对位置：'+loc+'。');p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(5)
        for r in p.runs:r.font.size=Pt(9)

with ZipFile(SOURCE) as z:
    source_comments={n:z.read(n) for n in z.namelist() if 'comments' in n.lower() and n.endswith('.xml')}
    cr=E.fromstring(z.read('word/comments.xml'))
    COMMENTS=[{'id':c.get(qn('w:id')),'author':c.get(qn('w:author')),'question':''.join(c.xpath('.//w:t/text()',namespaces=NS))} for c in cr]

D=Document(SOURCE)
for el in list(D._element.xpath('//w:commentRangeStart | //w:commentRangeEnd | //w:commentReference')):remove(el)
rewrite(find(D,'修订稿 ·'),'审阅修订稿 · 2026年10月24日')
rewrite(find(D,'本发明针对自然驾驶事件的交互扩展测试'), '本发明提供一种自然驾驶事件约束的自动驾驶多失效域测试方法。在固定道路、车辆身份、尺寸和起始状态的条件下，将背景行为表示为按车辆及时间索引的可替换参数组，并在实际交互使参考运动不再适用时执行响应控制。通过保持其余输入及随机实现一致的配对执行，利用单组变更前后的实际风险变化更新参数组选择概率；根据实际交互过程划分类别，在风险准入集合中联合保留类别代表并跨类别累计限制同一独立抽样来源的后续测试名额。按分配结果及所述概率生成下一轮仿真输入，局部变更继承来源标识，完整独立条件抽样新建来源标识，最终保存可重放的失效场景集合。')

def background(d):
    heading(d,'2 技术背景与相近的现有方案')
    para(d,'本节以自动驾驶闭环测试中的自然事件变体、交互引导生成及多样失效搜索为范围，选取十项直接相关的公开文献。各项文献用于说明具体技术基础；本方法的区别集中于实际执行反馈与后续测试输入分配之间的配合关系。')
    heading(d,'2.1 自然事件的响应式重建',2)
    para(d,'WO2025090341A1根据交通交互条件使背景车辆偏离日志轨迹，并考虑后续参与车辆的响应；MixSim通过路线条件行为策略构建自然场景的响应式变体。[1][2] 因而，保留自然事件依据并允许背景车辆动态响应，属于本案可以采用的执行基础。')
    para(d,'本方法进一步将实际响应产生的轨迹用于两项后续操作：核对单组行为输入变化的风险影响，以及形成用于测试名额分配的交互描述。二者使用同一执行结果，且与车辆及时间索引的行为输入相对应。')
    heading(d,'2.2 交互结构与测试反馈引导',2)
    para(d,'CCDiff及US20260134173A1根据交互因果结构识别关键车辆，对关键车辆的扩散生成施加目标引导。[3][4] CN121743096A依据语义嵌入和决策敏感区域调整场景参数。[8] 上述公开说明，利用交互关系或测试反馈选择干预对象已有技术基础。')
    para(d,'本方法采用可重放的单组变更比较：固定自然事件、当前策略、非选中行为输入及随机实现，比较原输入与变更输入经闭环执行后的实际风险，并据此分配下一轮各参数组的变更机会。车辆之间由交互引起的后续响应允许自然传播。')
    heading(d,'2.3 危险场景的多样性保留',2)
    para(d,'CaDRE在真实轨迹扰动空间中采用质量—多样性搜索和占用感知重启；STRIVE在学习交通先验中生成事故场景，并按碰撞类型组织结果。[5][6] CN116361175A将分类、组合测试与风险安全域划分相结合。[7] WO2025190506A2由车辆数据生成合理场景变体，并结合发生概率和严重程度生成测试覆盖；PtoP利用危险引导与粒子排斥获得分散的测试种子。[9][10]')
    para(d,'本方法把交互差异与抽样谱系作为两个独立维度：交互类别来自实际轨迹，来源标识来自完整独立条件抽样。局部变更即使进入其他类别仍沿用原来源；同一来源在所有类别中占用的名额合并计数。该联合约束直接决定后续仿真输入的生成机会，实际失效的保存另外执行。')
    heading(d,'2.4 本发明所针对的具体技术缺口',2)
    para(d,'危险程度、行为距离或类别覆盖能够描述样本的某一属性；同一初始抽样经反复局部变更也可能进入多个类别，因而类别分散并不能单独约束派生来源的集中。本方法在保持事件条件及输入索引一致的基础上，将参数组风险反馈、类别代表保留和跨类别来源容量连接至同一轮仿真输入安排，分别回答应改变何种行为、保留何种交互以及各独立来源获得多少后续机会。')
    refs(d)
replace_section(D,'2 技术背景','3 现有方案',background)

rewrite(find(D,'对于上述目标任务，'), '对于同一自然驾驶事件的定向测试，需要同时处理局部行为变更的有效性与测试群体的重复扩展。一次提议同时改变初态及多类行为输入时，风险差难以对应至具体行为组；按风险或类别分别保留候选时，同一抽样来源仍可能跨类别反复占用后续测试机会。')
rewrite(find(D,'在此基础上，本发明在风险准入集合内联合'), '在此基础上，本发明将风险准入候选的交互类别与独立来源建立关联，在同一次名额分配中保留类别代表并累计限制来源总名额；各名额随后调用由配对风险反馈确定的参数组提议规则。由此，类别与来源约束控制获得扩展机会的实例，风险反馈控制该实例中下一次直接变化的行为输入，共同作用于实际仿真调用。')
rewrite(find(D,'S101，'), 'S101，载入事件条件。从连续自然驾驶记录或已经构建的事件库取得目标事件，保存道路、车辆身份、尺寸、起始状态及有效历史。事件库可预先按风险筛选，本步骤输出所选事件的固定执行条件。')
after(find(D,'S106，'),lambda d:figure(d,1,'自然驾驶事件约束的多失效域测试流程','绘制S101事件载入、S102行为参数组、S103交互执行、S104配对风险反馈、S105联合名额分配、S106失效保存。S104输出参数组选择概率，S105输出父实例与名额，两者共同连接下一轮输入生成，再返回S103。独立条件抽样支路新建来源，局部变更支路继承来源；失效保存连接所有有效执行结果。'))

rewrite(find(D,'5.2 自然驾驶事件及风险尾部选择'),'5.2 自然驾驶事件载入与执行风险计算')
rewrite(find(D,'对自然事件风险超过拟合阈值'), '目标事件可由已经建立的关键事件库载入，也可按预设风险分位数或明确交互条件选取。采用极值尾部模型时，该模型在本轮开始前固定，作为事件选择接口；事件选择阈值与后续搜索门槛分别记录。进入同一事件的搜索后，道路、身份、尺寸、初态及有效历史保持一致。')
rewrite(find(D,'在一种实现中，条件归一化流'), '在一种实现中，条件归一化流根据初始条件和占位掩码生成各背景车辆的稀疏二维位置结点。以事件起点和这些结点进行保形分段三次Hermite插值，形成连续基准位置轨迹；条件扩散模型生成联合位置残差，残差经过端点及结点约束后与基准轨迹相加，得到背景车辆的软参考运动。车辆尺寸随事件保持，不作为轨迹残差的变化输入。其他能够保持固定事件条件、接收显式行为输入并重建参考的模型，也可实现该生成接口。')
after(find(D,'接管标记与实际运动改变分别记录。'),lambda d:figure(d,2,'参考适用性判断、持续接管与机动承接','绘制实际联合状态、参考跟踪器、参考适用性判断、持续接管状态、背景响应策略及统一动力学。触发时继承当前换道目标、进度和上一动作；分别标注控制权已接管与实际动作已改变，实际运动反馈至下一步关系判断。图中不预设接管后必然避免碰撞。'))
after(find(D,'上述规则用于定向发现失效。'),lambda d:figure(d,3,'同条件参数组比较与后续提议','并列绘制原输入与只替换一个参数组的输入；共享事件、策略版本、其余参数及随机数组。两条实际轨迹进入风险差计算，再更新敏感度及选择概率。标出车辆掩码、作用时间范围与未选中元素保留，并将输出连接到图1的下一轮输入生成。'))
rewrite(find(D,'式（8）中，'),r'式（8）中，$s_b$为累计敏感度，$\rho$为介于0与1之间的更新权重，$a_b$为交互相关权重，$\delta$为正平滑常数；$b^{\prime}$为参数组求和索引。输出$q_b$为后续局部提议选择参数组$b$的概率。尚无有效配对时保留预设初始敏感度，使适用但未充分探索的组仍有机会被选择。')
rewrite(find(D,'按分配结果使用第5.6节提议规则'), '联合分配输出包含父实例、所属类别、继承来源及提议次数的测试安排。每个名额按第5.6节的参数组选择概率产生局部变更，经同一执行模块计算风险，并按当前门槛接纳或拒绝；子实例重新计算交互类别，但继承来源。完整独立条件抽样使用新的行为输入并新建来源。执行结果继续更新敏感度、类别描述及下一轮容量，形成逐轮可核对的输入—执行—分配关系。被拒绝提议的成本和实际失效仍被保存。图4说明联合约束。')
rewrite(find(D,'本发明的测试输出不依赖特定更新算法'), '失效档案通过固定的实例恢复接口向后续开发流程输出测试输入。策略更新作为可选用途；更新后重新执行测试并记录策略版本。')
after(find(D,'联合分配输出包含'),lambda d:figure(d,4,'交互类别与独立来源的联合名额分配','绘制风险准入候选到类别—来源二部图的映射：类别侧代表容量为1，来源侧为跨类别共享的剩余名额上限。最大流产生代表后在剩余容量内填充；输出父实例及提议次数，连接参数组选择与下一轮仿真。用同一来源的连线跨越不同类别表示统一累计，容量不足进入独立抽样或缩减实际名额分支。本图采用结构示意，不填入虚构实验数值。'))

def protect(d):
    heading(d,'6 关键技术特征与欲保护内容')
    heading(d,'6.1 独立方法保护的技术关系',2)
    para(d,'建议将以下相互依赖的处理关系共同作为方法保护的核心。其一，取得自然驾驶事件并固定道路、车辆身份、尺寸、起始状态及有效历史，将允许变化的背景行为表示为可逐组替换和重放的输入；完整独立条件抽样建立来源标识，局部派生保留该标识。其二，依据实际交通状态执行参考跟踪及触发后的背景响应，在承接当前运动与已开始机动的条件下得到实际轨迹、风险和交互描述。')
    para(d,'其三，保持固定事件、当前策略、非选中行为输入及随机实现一致，对单组变更前后的实际执行结果进行比较，以风险变化更新后续参数组选择概率。其四，对风险准入候选联合保留交互类别代表，并对同一来源在全部类别中的后续名额求和及限额；派生实例改变类别时继续沿用原来源。其五，依据所分配名额选择父实例，并按所述选择概率生成下一轮行为输入，执行结果返回风险比较与联合分配，失效实例连同完整输入单独保存。')
    para(d,'上述关系把“何种实例获得扩展机会”和“该实例的何种行为输入被改变”连接至同一仿真调用。固定事件条件为风险配对提供一致背景；实际轨迹同时形成风险反馈与交互描述；来源继承使跨类别容量约束在后续派生中持续有效。')
    heading(d,'6.2 从属保护的具体限定',2)
    para(d,'第一组限定涉及可重放输入：按稳定车辆身份、物理帧及变量类型预先索引随机数组，未选中元素的索引和值保持一致；从事件起点重放并恢复控制器、机动与策略状态。')
    para(d,'第二组限定涉及交互执行：依据短时冲突预测、经验误差支持或已标定交互刺激判断参考适用性；接管后保留控制权，分别记录接管标记与实际动作变化；从实际位置及既有目标继续已经开始的换道。采用空间路径查询的实施方式，可按实际纵向位置或道路弧长重查参考。')
    para(d,'第三组限定涉及变更选择：只使用双方均有效的配对结果，以归一化且截断的绝对风险差更新敏感度，结合交互相关范围与正平滑项确定参数组概率；局部提议与完整独立条件抽样分别处理来源继承及新建。')
    para(d,'第四组限定涉及联合分配：建立类别—来源二部图，以类别代表容量和来源共享上限计算代表分配，再于剩余容量内填充名额；容量不足时补充经相同门槛准入的独立抽样，或减少实际名额，保持来源累计规则。')
    para(d,'第五组限定涉及存储及成本：完整失效记录与容量受限的代表档案分开保存；为输入、版本和轨迹建立内容核验信息；配对、拒绝、独立补充及异常执行均计入实际仿真消耗。')
    heading(d,'6.3 系统与设备保护',2)
    para(d,'相应系统包括事件处理模块、行为输入生成模块、交互执行模块、风险比较模块、联合名额分配模块和失效存储模块。事件处理模块输出固定条件；行为输入生成模块接收父实例、名额及参数组选择概率；交互执行模块输出实际轨迹；风险比较模块输出风险准入结果及更新后的选择概率；联合分配模块输出下一轮父实例与名额；失效存储模块保存输入、执行结果及谱系。处理器和存储器执行上述模块之间的数据传递与计算，构成对应测试设备及程序实现。')
replace_section(D,'6 关键技术特征','7 有益效果',protect)

def examples(d):
    heading(d,'8 具体实施例与已有工程材料')
    heading(d,'8.1 高速公路自然事件与参考生成',2)
    para(d,'以highD高速公路轨迹为输入，将车辆中心、速度及车道信息转换为统一道路坐标。窗口保留25 Hz采样和150个状态；在起点选择同车道及相邻车道最近前后车辆，最多保留六个背景位置，稳定身份贯穿整个窗口。现有事件库含96,055个合格片段，条件生成资产使用2,964个目标尾部事件。上述数据用于给出可实施的数据规模与接口。')
    p=para(d,'数据来源：Krajewski R, Bock J, Kloeker L, Eckstein L. The highD Dataset，ITSC，2018。');link(p,'原文','https://arxiv.org/abs/1810.05642')
    para(d,'将所选事件的初始条件、占位掩码及车辆尺寸传入条件采样接口；初始条件为26维，六个背景位置分别在2.0、4.0和5.96 s设置二维位置结点。固定事件比较只改变结点基输入、残差生成噪声、驾驶员及换道参数和响应随机数组。图5用于展示自然事件、条件参考与实际执行之间的关系。')
    figure(d,5,'自然事件载入与条件参考生成的工程示例','从既有事件库与条件生成输出中选取一个完整可重建事件，绘制道路、稳定车辆编号及起点包围框；并列展示原记录、稀疏结点与生成软参考，固定初态、占位和尺寸采用相同标记。图注记录录像、车辆及窗口标识和模型版本。该图展示条件生成组件；接入本案完整固定事件搜索后，再补充两次执行的固定字段一致性核验，不将一般分布生成结果标为固定事件配对结果。')
    heading(d,'8.2 背景响应与机动承接',2)
    para(d,'交通世界按0.04 s物理步长执行，背景车辆每步检查实际状态和参考适用性。现有配置使用1.5 s短时预测窗及1/16冲突支持门槛，并结合跟驰和前车离开刺激。背景车辆接管时承接当前换道及实际位置，分别记录控制权状态、实际动作变化和触发原因。')
    para(d,'既有自然记录重建报告涵盖10,151个不同场景；控制干预材料包含51个预选案例和204个回合。其主车采用脚本控制，适于展示响应机制。选取已保存的制动、前车离开及相应对照案例，可据图6检查交互刺激、接管时刻与实际运动之间的对应关系。')
    figure(d,6,'既有控制干预案例的背景响应过程','优先从已保存回放清单中选取制动案例1及其无干预对照案例0，必要时补充前车离开案例3；使用同一随机种子和原有输出。上部为道路俯视关键帧，下部为车辆间距、纵向速度、跟踪动作、实际动作及接管标记的时间曲线。标出实际记录的触发时刻，碰撞发生时照实保留；完整绘图数据不足时重放该既有案例补齐。图注注明脚本干预与响应组件验证的用途。')
    para(d,'现有CPU计时材料记录最多七个参与实体的8,493个物理步，世界更新耗时99%分位约8.59 ms，未记录超过40 ms的步骤；计时不含被测策略计算。该记录说明当前执行组件在所测配置下的计算开销。')
    heading(d,'8.3 风险反馈、交互描述与来源记录',2)
    para(d,'现有搜索记录已经保存父实例、独立来源、变更组、实际风险、关键车辆、关键时刻及交互描述，并依据单组提议前后的风险差更新敏感度。实施本案时，先核对事件固定字段与非选中输入一致，再使用第5.6节的有效配对更新规则；随后按第5.7节从实际轨迹形成类别，将类别及继承来源共同送入第5.8节的名额分配器。')
    figure(d,7,'实际搜索记录中的风险反馈与交互描述','从已有搜索日志中筛选具有父实例、单一变更组且双方执行有效的配对，关联实例文件，展示参数组、车辆及时间范围、前后风险和关键交互状态。现有一般生成搜索配对标注为组件输出；只有核对道路、初态、身份、尺寸和非选中数组一致的配对，才能标注为本案固定事件比较。联合分配接入后，用真实分配日志补充类别—来源连接、累计占用和下一轮父实例，不再采用手工名额数值算例。')
    heading(d,'8.4 完整实施的核验与失效复现',2)
    para(d,'测试设备按照S101—S106连接事件读取、条件行为生成、背景响应、风险比较、联合分配及失效存储。通过输入一致性核验确认固定字段，逐条核对局部派生的来源继承；对每轮分配，分别核对风险准入、类别代表、跨类别来源累计及名额不足处理，并核对分配结果确实进入后续仿真调用。')
    figure(d,8,'联合分配后的实际测试集合与失效复现证据','完整接通后，按同一策略、事件池、物理配置及实际执行预算整理结果。左侧绘制实际类别—来源名额矩阵，右侧展示若干不同交互失效的实际轨迹及复现对照；需要比较发现效果时，分别加入无来源限额和无参数组风险反馈的同预算结果。未取得联合分配日志前保留占位，不绘制预测增益曲线。')
    para(d,'失效复现时恢复事件、全部行为输入或参考重建输入、策略及环境版本和随机实现，从事件起点执行，核对输入哈希、关键风险时刻及失效标签。完整记录持续保留，代表档案的容量变化不影响已经发现失效的重放。')
    para(d,'本节以已存在的工程记录说明可运行组件，并明确完整实施的核验对象。固定事件搜索接入、跨类别来源容量分配及其联动效果的工程核对，见随附《新颖性创造性与前案区别分析及审阅意见处理说明》。')
replace_section(D,'8 具体实施例',None,examples)

# The symbol registry is the sole substantive table outside the cover.
registry=D.tables[1]
for row in [
('帧索引／车辆索引','$t,T,i$','物理帧、末帧及稳定背景车辆索引；预测时间点集合另记为$\\mathcal{T}$。'),
('参数组选择与更新','$b,s_b,q_b$','参数组索引、累计敏感度及组选择概率；交互权重$a_b$与正平滑项$\\delta$配合。'),
('提议及独立抽样','$K_{e,b},K_e,\\mu_e$','局部提议、混合提议及完整条件抽样律；独立比例为$\\varepsilon$。'),
('分层得分／保留比例','$Z_n,p_0$','有效候选$n$的有界分层得分及风险层目标保留比例。'),
('类别集合／候选名额','$\\mathcal G_\\ell,m_n$','本层准入集合中的交互类别集合，以及候选获得的后续名额。'),
('来源名额／上限','$M_\\alpha,c_\\alpha$','同一来源跨全部类别的名额总量，以及该来源的名额上限。'),
('有效来源数','$A_{\\mathrm{eff}}$','由本轮名额占比计算的来源分散度；统一来源上限记为$c$。')]:
    cells=registry.add_row().cells
    for cell,value in zip(cells,row):
        cell.text='';rich(cell.paragraphs[0],value,10)
        for sz in cell._tc.xpath('.//m:r/w:rPr/w:sz'):sz.set(qn('w:val'),'20')
        cell.paragraphs[0].paragraph_format.line_spacing=1.1
        cell.paragraphs[0].paragraph_format.space_after=Pt(4)
    cp=OxmlElement('w:cantSplit');registry.rows[-1]._tr.get_or_add_trPr().append(cp)

# Render native OMML equations in borderless three-cell rows, with aligned right numbers.
equations=[]
for p in list(D.paragraphs):
    mp=p._p.find(qn('m:oMathPara'))
    match=re.fullmatch(r'\s*（(\d+)）\s*',p.text)
    if not match:continue
    if mp is None:
        om=p._p.find(qn('m:oMath'))
        if om is None:continue
        mp=OxmlElement('m:oMathPara');mp.append(deepcopy(om))
    number=match.group(1)
    if number=='8':
        mp=OxmlElement('m:oMathPara')
        mp.append(math(r's_b\leftarrow(1-\rho)s_b+\rho d_b,\quad q_b=\frac{a_b s_b+\delta}{\sum_{b^{\prime}\in\mathcal B_e}(a_{b^{\prime}}s_{b^{\prime}}+\delta)}'))
    tb=D.add_table(rows=1,cols=3);tb.autofit=False
    widths=[1.4,13.4,1.4]
    for c,w in zip(tb.columns,widths):c.width=Cm(w)
    pr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ('top','left','bottom','right'):
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for idx,c in enumerate(tb.rows[0].cells):
        c.width=Cm(widths[idx]);c.vertical_alignment=1
        cp=c.paragraphs[0];cp.paragraph_format.space_before=Pt(4);cp.paragraph_format.space_after=Pt(4)
        cp.paragraph_format.keep_with_next=True
        cp.paragraph_format.first_line_indent=Cm(0)
    cp=tb.cell(0,1).paragraphs[0];cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp._p.append(deepcopy(mp))
    cp=tb.cell(0,2).paragraphs[0];cp.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(cp.add_run('（'+number+'）'),10)
    cp=OxmlElement('w:cantSplit');tb.rows[0]._tr.get_or_add_trPr().append(cp)
    # Keep the equation and its symbol explanation on the same page.
    explanation=p._p.getnext()
    assert explanation.tag==qn('w:p') and txt(explanation).startswith('式（'+number+'）')
    wrap=D.add_table(rows=1,cols=1);wrap.autofit=False;wrap.columns[0].width=Cm(16.2)
    wrap._tbl.tblPr.append(deepcopy(borders));wrap._tbl.tblPr.append(deepcopy(margins))
    cell=wrap.cell(0,0);cell.width=Cm(16.2)
    for x in list(cell._tc):
        if x.tag!=qn('w:tcPr'):cell._tc.remove(x)
    cell._tc.append(tb._tbl);cell._tc.append(deepcopy(explanation))
    cp=OxmlElement('w:cantSplit');wrap.rows[0]._tr.get_or_add_trPr().append(cp)
    p._p.addprevious(wrap._tbl);remove(explanation);remove(p._p);equations.append(int(number))
assert equations==list(range(1,15)), equations
for omrun in D._element.xpath('//m:r'):
    vals=[x.text for x in omrun.findall(qn('m:t'))]
    if vals and vals[0] in ('min','max','clip','exp'):
        pr=omrun.find(qn('m:rPr'))
        if pr is None:pr=OxmlElement('m:rPr');omrun.insert(0,pr)
        for sty in list(pr.findall(qn('m:sty'))):pr.remove(sty)
        sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty)

# Keep all original reviewer texts, authors, IDs and modern metadata; re-anchor only.
ANCHORS={
'0':'2 技术背景与相近的现有方案',
'1':'5.1 数据对象、固定条件与重放约定',
'2':'每组行为输入同时记录所属车辆',
'3':'6 关键技术特征与欲保护内容',
'4':'8 具体实施例与已有工程材料',
'5':'本节以已存在的工程记录',
'6':'图1  自然驾驶事件约束的多失效域测试流程',
'7':'2.5 本节引用的相近公开文献'}
for cid,prefix in ANCHORS.items():
    p=find(D,prefix);st=OxmlElement('w:commentRangeStart');st.set(qn('w:id'),cid)
    en=OxmlElement('w:commentRangeEnd');en.set(qn('w:id'),cid)
    pos=1 if p._p.find(qn('w:pPr')) is not None else 0;p._p.insert(pos,st);p._p.append(en)
    run=OxmlElement('w:r');ref=OxmlElement('w:commentReference');ref.set(qn('w:id'),cid);run.append(ref);p._p.append(run)

main=OUT/'技术交底书_自然驾驶事件约束的自动驾驶多失效域测试方法_审阅修订稿.docx'
D.core_properties.title='自然驾驶事件约束的自动驾驶多失效域测试方法（审阅修订稿）'
D.save(main)
# Ensure metadata bytes remain exactly those authored by Word in the input.
with ZipFile(main) as z:data={n:z.read(n) for n in z.namelist()}
for n,b in source_comments.items():
    assert n in data,n
    data[n]=b
with ZipFile(main,'w',compression=8) as z:
    for n,b in data.items():z.writestr(n,b)

RESPONSES=[
('2及2.5','已按自然事件响应、交互引导与多样性搜索选取十项直接相关文献，在第2节分析并就地列出原文链接和核对位置；补充WO2025190506A2及PtoP，移出通用子集模拟文献。'),
('5.1及各公式释义','保留并扩展符号表，统一事件、行为输入、参数组、交互类别、独立来源和名额；预测时间点集合与末帧索引分别表示。全部十四个公式保留原生Word数学对象，逐式保留输入、输出和符号释义。'),
('公式（1）—（14）','十四个公式统一使用无边框三列布局：中列居中公式、右列右对齐编号，编号顺序连续。公式及编号分开保存，仍可编辑。'),
('6.1—6.3','删除该节表格，将核心技术关系、五组从属保护特征和系统模块改为连续技术段落，补足输出进入下一轮仿真输入的关系。'),
('8.1—8.4，图5—图8','删除十二名额的手工数值算例及其表格；使用已跑通组件、实际报告和日志描述实施过程，加入可追溯的图片占位及所需数据说明；区分脚本干预、一般生成搜索与完整联合实施。'),
('8末段及第5节对应步骤','删除原第9节及其泛化应用描述；适用且已有支持的生成接口、类别定义和分配求解选择放回相应实施步骤，不新增道路或参与者实施效果。'),
('4、5.4、5.6、5.8及8','删除原独立附图章节；图1—图4移入各自技术步骤，图5—图8置于对应实施例，图名、占位、用途和绘制要求就地说明。'),
('2.5及8.1','删除文末独立参考文献节。十项相近技术文献置于第2节；highD作为数据来源就近置于8.1，不混作最接近技术方案。')]
for c,(loc,response) in zip(COMMENTS,RESPONSES):c.update(location=loc,response=response,anchor_now=ANCHORS[c['id']],status='已落实正文修改')
(OUT/'support/review_response_map.json').write_text(json.dumps(COMMENTS,ensure_ascii=False,indent=2))
(OUT/'support/references.json').write_text(json.dumps(REFS,ensure_ascii=False,indent=2))
print(main)
