"""Validate the reconceived patent draft against its text and existing evidence."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import hashlib
import json
import re
import sys

sys.path.insert(0, '/tmp/patent_drafting_deps')
from docx import Document

ROOT = Path(__file__).resolve().parents[4]
OUT = Path(__file__).resolve().parents[1]
DOC = OUT / '技术交底书_自动驾驶扩散规划器的交互失效修正方法.docx'
MD = DOC.with_suffix('.md')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}

with ZipFile(DOC) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith(('.xml', '.rels')):
            etree.fromstring(z.read(name))
    xml = etree.fromstring(z.read('word/document.xml'))
    settings = etree.fromstring(z.read('word/settings.xml'))
    formulas = xml.findall('.//m:oMath', NS)
    inline = len(re.findall(r'(?<!\$)\$[^$]+\$(?!\$)', MD.read_text()))
    display = len(re.findall(r'^\$\$', MD.read_text(), re.M))
    assert display == 8
    assert len(formulas) == inline + display, (len(formulas), inline, display)
    for f in formulas:
        s = ''.join(f.xpath('.//m:t/text()', namespaces=NS))
        assert not re.search(r'[\u4e00-\u9fff，。；：]', s), s
    assert not xml.xpath('//w:pict | //w:object | //w:del | //w:ins | //w:commentRangeStart', namespaces=NS)
    assert not xml.xpath('//w:instrText[contains(.,"EQ ")]', namespaces=NS)
    normal = ''.join(xml.xpath('//w:t/text()', namespaces=NS))
    assert not re.search(r'\$|\\(?:frac|omega|pi|tau)|[_^]', normal)
    numbers = [s for s in xml.xpath('//w:t/text()', namespaces=NS) if re.fullmatch(r'（\d+）', s)]
    assert numbers == ['（' + str(i) + '）' for i in range(1, 9)], numbers
    assert settings.xpath('//w:compatSetting[@w:name="compatibilityMode"]/@w:val', namespaces=NS) == ['15']
    assert not [n for n in z.namelist() if n.startswith('word/media/')]
    headings = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading1"]]', namespaces=NS)
    subheadings = xml.xpath('//w:p[w:pPr/w:pStyle[@w:val="Heading2"]]', namespaces=NS)
    assert len(headings) == 7 and len(subheadings) == 4
    assert normal.count('（附图占位）') == 4
    refs = re.findall(r'^\[(\d+)\]', MD.read_text(), re.M)
    assert refs == [str(n) for n in range(1, 7)]
    rels = etree.fromstring(z.read('word/_rels/document.xml.rels'))
    urls = [r.get('Target') for r in rels if r.get('Type').endswith('/hyperlink')]
    expected_urls = [re.search(r'https?://\S+$', line).group(0)
                     for line in MD.read_text().splitlines()
                     if re.match(r'^\[\d+\]', line)]
    assert urls == expected_urls, urls
    assert all(url.startswith('https://') for url in urls)
    for line in MD.read_text().splitlines():
        if re.match(r'^\[\d+\]', line):
            citation = line[:re.search(r'https?://\S+$', line).start()].rstrip()
            assert citation in normal, citation
    main_body = MD.read_text().split('## 7 ')[0]
    assert set(re.findall(r'\[(\d+)\]', main_body)) == set(refs)
    for removed_reference in ['SERA', 'WO2019232335', 'US20230227069', 'Learning without Forgetting']:
        assert removed_reference not in MD.read_text()
    assert '不进入交通观测' in normal
    assert '功能误差约束模型输出变化' in normal
    assert '更新投影' not in normal and '偏移状态补标' not in normal

readback = Document(DOC)
assert len(readback.tables) == 1
assert readback.core_properties.title == '一种自动驾驶扩散规划器的交互失效修正方法'
assert not list(OUT.glob('*.pdf')) and not list(OUT.glob('*.svg'))
abstract = MD.read_text().split('**摘要建议稿**')[1].strip()
assert len(abstract) < 300

claims = (OUT / '候选权利要求与说明书支持.md').read_text()
claim_numbers = re.findall(r'^\*\*(\d+)\.\*\*', claims, re.M)
assert claim_numbers == [str(n) for n in range(1, 10)]
assert '权利要求1至7' in claims

manifest = json.loads((OUT / 'support/source_manifest.json').read_text())
for item in manifest:
    path = ROOT / item['path']
    assert hashlib.sha256(path.read_bytes()).hexdigest() == item['sha256'], item['path']

base = ROOT / 'results/risk_stratified_repair/formal/diffusion_planner/preparation'
materials = json.loads((base / 'round_one_material_acceptance.json').read_text())['methods']['guided']
assert materials['material_inputs'] == 64
assert materials['expert_inputs'] == 27 and materials['policy_success_inputs'] == 37
assert materials['new_defect_or_neighbor_inputs'] == 32
replay = json.loads((base / 'guided_round_one_label_replay.json').read_text())
assert replay['completed_selected_expert_caches'] == 27
assert replay['total_unique_diagnostic_executions'] == 27
assert replay['policy_success_inputs'] == 5
assert all(replay[k] == 0 for k in ['collisions', 'offroad', 'stalled'])
assert replay['scope'] == 'Offline oracle label tracking; not ADS or final evaluation'
admission = json.loads((base / 'guided_round_one_admission_acceptance.json').read_text())
assert admission['selected_candidate_updates'] == 500
candidates = admission['candidates']
assert [c['updates'] for c in candidates] == [500, 1000]
for c, expected in zip(candidates, [(15, 9), (50, 27)]):
    comparisons = c['comparisons']
    tail, natural = comparisons['official/tail_generated'], comparisons['official/natural_generated']
    assert tail['instances'] == natural['instances'] == 500
    actual = (round(tail['collision_rate_after'] * 500), round(natural['collision_rate_after'] * 500))
    assert actual == expected
    assert round(tail['collision_rate_before'] * 500) == 20
    assert round(natural['collision_rate_before'] * 500) == 10
assert admission['full_checkpoint_payload_equal_to_selected']

report = {
    'date': '2026-10-25',
    'revision': 'focused_review',
    'structural_checks_passed': True,
    'native_equations': len(formulas),
    'inline_equations': inline,
    'numbered_equations': display,
    'heading_1': len(headings),
    'heading_2': len(subheadings),
    'logic_figure_placeholders': 4,
    'figures_generated': 0,
    'claims_checked': len(claim_numbers),
    'direct_prior_art_comparisons': 5,
    'implementation_baseline_references': 1,
    'reference_links_checked': len(urls),
    'sources_hash_checked': len(manifest),
    'new_training_simulation_or_expert_calls': 0,
    'word_desktop_reopen': False,
    'word_note': 'No desktop Word session is available. ZIP/XML, native Office Math and python-docx readback checked. LibreOffice rendering is used for layout QA.',
    'abstract_characters': len(abstract),
    'docx_sha256': hashlib.sha256(DOC.read_bytes()).hexdigest(),
    'pdf_delivered': False,
    'svg_delivered': False,
    'logic_review': '已复核样本来源、因果观测、真实未来、去噪输入、目标选择与闭环接纳；功能约束与行为验证分别说明。',
    'prior_art_review': '五项直接对照及一项实施模型基础；US20250171046A1已纳入，计划与实际执行的区别不单独主张为创新。',
    'experimental_scope': '仅引用既有完整流程结果，无新增消融或实验，不将整体效果归因于单项机制。',
    'previous_draft_preserved': 'support/before_focused_review/',
}
path = OUT / 'support/qa_report.json'
if path.exists():
    previous = json.loads(path.read_text())
    if previous.get('docx_sha256') == report['docx_sha256']:
        for key in ['layout_pages', 'layout_review', 'equation_number_alignment']:
            if key in previous:
                report[key] = previous[key]
path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False))
